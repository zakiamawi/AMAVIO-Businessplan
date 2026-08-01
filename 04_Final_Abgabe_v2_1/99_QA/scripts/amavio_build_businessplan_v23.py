from __future__ import annotations

import json
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


NAVY = "1E3A5F"
INK = "0F1B2E"
BRASS = "C9A961"
CREAM = "F4F1EA"
WARM = "F8F5EE"
COPPER = "B97855"
PLATINUM = "D8D4C8"


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "01_Businessplan" / "AMAVIO_Living_Businessplan_Master_2026_v2_2_FINAL.docx"
OUT = ROOT / "01_Businessplan" / "AMAVIO_Living_Businessplan_Master_2026_v2_3_AUSFUEHRLICH_ARBEITSSTAND.docx"
QA = ROOT / "99_QA" / "AMAVIO_Businessplan_v2_3_Ausfuehrlich_QA_20260731.json"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_text(cell, text: str, bold: bool = False, color: str = INK, size: float = 9.5) -> None:
    cell.text = ""
    para = cell.paragraphs[0]
    para.paragraph_format.space_after = Pt(2)
    run = para.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_paragraph(paragraph, size: float = 10.5, color: str = INK, bold: bool = False) -> None:
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.08
    for run in paragraph.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(color)
        run.bold = bold


def add_heading(doc: Document, text: str, level: int = 1, color: str = NAVY) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Georgia"
    run.font.size = Pt(18 if level == 1 else 13.5)
    run.font.color.rgb = RGBColor.from_string(color)


def add_body(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(text)
    style_paragraph(paragraph)


def add_bullets(doc: Document, items: list[str]) -> None:
    rows = [[str(idx), item] for idx, item in enumerate(items, start=1)]
    add_table(doc, ["#", "Prüfpunkt"], rows)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        set_cell_shading(hdr[i], NAVY)
        set_cell_text(hdr[i], header, bold=True, color="FFFFFF", size=8.8)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_shading(cells[i], WARM if len(table.rows) % 2 == 0 else "FFFFFF")
            set_cell_text(cells[i], value, color=INK, size=8.8)
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            mar = tc_pr.find(qn("w:tcMar"))
            if mar is None:
                mar = OxmlElement("w:tcMar")
                tc_pr.append(mar)
            for side, value in {"top": "90", "left": "120", "bottom": "90", "right": "120"}.items():
                elem = mar.find(qn(f"w:{side}"))
                if elem is None:
                    elem = OxmlElement(f"w:{side}")
                    mar.append(elem)
                elem.set(qn("w:w"), value)
                elem.set(qn("w:type"), "dxa")
    doc.add_paragraph()


def add_callout(doc: Document, title: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = table.rows[0].cells
    set_cell_shading(left, BRASS)
    set_cell_shading(right, CREAM)
    set_cell_text(left, title.upper(), bold=True, color=INK, size=8.5)
    set_cell_text(right, text, bold=False, color=INK, size=9.5)
    doc.add_paragraph()


def replace_version_in_container(container) -> None:
    for paragraph in container.paragraphs:
        for run in paragraph.runs:
            if "v2.2" in run.text:
                run.text = run.text.replace("v2.2", "v2.3 Arbeitsstand")
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                replace_version_in_container(cell)


def update_footers(doc: Document) -> None:
    for section in doc.sections:
        replace_version_in_container(section.footer)
        replace_version_in_container(section.first_page_footer)
        replace_version_in_container(section.even_page_footer)


def append_readiness_section(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "12 · PRÜFER- UND COACH-READINESS-ERGÄNZUNG", 1)
    add_body(
        doc,
        "Diese Ergänzung gehört zur ausführlichen Arbeitsfassung v2.3. Sie verändert den bestehenden Businessplan-Kern nicht, sondern macht die prüfungsrelevante Leselogik für Coach, Finanzierungsgespräch und finale Abgabe explizit sichtbar.",
    )
    add_callout(
        doc,
        "Führungslinie",
        "Führend bleibt der Businessplan v2.2 mit Finanzmodell v1.8. Diese v2.3-Arbeitsfassung fügt eine prüferorientierte Brücke hinzu und darf erst nach Sicht-/Rendercheck zur finalen Version erhoben werden.",
    )

    add_heading(doc, "12.1 Executive Prüflogik", 2)
    add_body(
        doc,
        "Der Businessplan soll nicht nur die Attraktivität der Idee zeigen, sondern beweisen, dass AMAVIO mit begrenztem Startkapital ein reales B2B-Problem in eine kontrollierbare Dienstleistungslogik überführt. Die stärkste Erzählung ist deshalb nicht Wachstum um jeden Preis, sondern Wachstum nach Beleg.",
    )
    add_bullets(
        doc,
        [
            "Problem: Internationale Mitarbeitende starten nicht verlässlich, wenn Wohnraum, Bezugsfertigkeit und Ankunftsorganisation ungeordnet bleiben.",
            "Lösung: AMAVIO bündelt Bedarfsklärung, Housing-Projektsteuerung, Bezugsfertigkeit, Arrival Office und Reporting in einem verantworteten Prozess.",
            "Startmodell: kapitalarmes Projekt- und Servicemodell vor eigenem Objektbetrieb.",
            "Proof-Logik: bezahltes Mandat, dokumentierte Delivery, wiederholbarer Playbook-Standard und erst danach höhere Kapitalbindung.",
        ],
    )

    add_heading(doc, "12.2 Führende Finanzlinie v1.8", 2)
    add_table(
        doc,
        ["Kennzahl", "Jahr 1", "Jahr 2", "Jahr 3", "Prüferlesart"],
        [
            ["Umsatz", "93.500 EUR", "269.000 EUR", "442.000 EUR", "Menge x Preis x Zeitpunkt x Kapazität"],
            ["EBITDA", "34.950 EUR", "120.320 EUR", "169.300 EUR", "Dienstleistungsmodell mit kontrollierter Fixkostenbasis"],
            ["Startmittel", "25.000 EUR", "-", "-", "500 EUR Stammkapital plus 24.500 EUR dokumentierte Zuführung/Darlehen"],
        ],
    )
    add_body(
        doc,
        "Im Gespräch sollte die Finanzplanung nicht als pauschaler Wachstumspfad verteidigt werden. Belastbarer ist die Herleitung über Angebotsmix, Zahlungszeitpunkte, Vorleistungen, Kapazitätsgrenzen und konkrete Mandatswahrscheinlichkeit.",
    )

    add_heading(doc, "12.3 P1-Gates vor finaler Abgabe", 2)
    add_table(
        doc,
        ["Gate", "Warum kritisch", "Erforderlicher Nachweis", "Statuslogik"],
        [
            ["Mittelzufluss 24.500 EUR", "Finanzierung darf nicht nur behauptet sein.", "Beleg, Darlehens-/Einlagezuordnung, Zahlungsweg", "Vor Abgabe klären"],
            ["§34c-/Vermittlungsgrenze", "Housing- und Objektpartnerlogik kann erlaubnispflichtige Bereiche berühren.", "saubere Rollenabgrenzung und ggf. juristische Einordnung", "nicht offensiv versprechen"],
            ["Unternehmensgegenstand", "Leistungsbild muss zum Gesellschaftszweck passen.", "Abgleich Satzung/Leistungsportfolio", "vor Coachtermin prüfen"],
            ["Nachfragebeweis", "Pipeline darf nicht als gesicherter Umsatz erscheinen.", "Zielaccountliste, Gesprächsstand, nächster Schritt", "als Proof-Pipeline kennzeichnen"],
            ["Partnerfähigkeit", "Lieferfähigkeit darf nicht wie feste Kapazität wirken.", "Partnerprofil, Rolle, Kapazität, Abnahmeprozess", "nur belastbar bestätigte Partner führen"],
        ],
    )

    add_heading(doc, "12.4 Ideales erstes Mandat", 2)
    add_body(
        doc,
        "Das ideale Erstmandat ist kein beliebiger Objektauftrag, sondern ein klar abgegrenzter Cohort Landing Launch. Ein Arbeitgeber oder Recruitingpartner hat eine konkret startende Gruppe, eine klare Zeitachse, Entscheidungsfähigkeit und die Bereitschaft, für koordinierte Umsetzung zu zahlen.",
    )
    add_table(
        doc,
        ["Kriterium", "Zielbild", "Nicht geeignet"],
        [
            ["Kohorte", "5 bis 20 internationale Mitarbeitende mit konkretem Startfenster", "unklare Interessenten ohne Startdatum"],
            ["Entscheider", "HR/Projektverantwortlicher mit Budgetzugriff", "reine Informationsanfrage ohne Mandatsmacht"],
            ["Scope", "Bedarf, Wohnoptionen, Bezugsfertigkeit, Ankunft und Reporting", "lose Wohnungssuche ohne Prozessverantwortung"],
            ["Zahlung", "Blueprint, Launch oder Arrival Office als bezahlte Leistung", "unbezahlter Suchauftrag als Akquiseersatz"],
        ],
    )

    add_heading(doc, "12.5 Abschlussbewertung für die nächste Bearbeitung", 2)
    add_body(
        doc,
        "Die ausführliche Fassung ist inhaltlich tragfähig, sollte aber vor finaler Abgabe nicht weiter mit globalen Designexperimenten belastet werden. Der nächste qualitative Sprung entsteht durch selektive Veredelung: Tabellen nur dort kompakter machen, wo sie wirklich stören; wichtige Prüflogiken als Callouts stärken; keine grünen Akzente; Navy, Creme, Brass und optional Copper als dominante Identität halten.",
    )
    add_bullets(
        doc,
        [
            "Keine bestehenden FINAL-Dateien ohne Backup überschreiben.",
            "Vor finalem Export Businessplan rendern und kritische Tabellen-/Seitenbrüche visuell prüfen.",
            "Decks und Website erst nach stabiler Businessplan-Linie weiter visuell eskalieren.",
        ],
    )


def main() -> int:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if OUT.exists():
        raise FileExistsError(f"Refusing to overwrite existing output: {OUT}")

    shutil.copy2(SOURCE, OUT)
    doc = Document(OUT)
    before = {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
    }
    append_readiness_section(doc)
    update_footers(doc)

    props = doc.core_properties
    props.title = "AMAVIO Living Businessplan Master 2026 v2.3 - Ausführliche Arbeitsfassung"
    props.subject = "Workforce Housing, Arrival Management und Integration"
    props.comments = "Ausführliche Arbeitsfassung auf Basis v2.2 FINAL plus Prüfer-/Coach-Readiness-Ergänzung vom 31. Juli 2026."
    props.keywords = "AMAVIO, Businessplan, Workforce Housing, Finanzmodell v1.8, Coach Readiness"
    props.modified = datetime(2026, 7, 31, 12, 0, 0)

    doc.save(OUT)

    checked = Document(OUT)
    full_text = "\n".join(p.text for p in checked.paragraphs)
    needles = ["v1.7", "v1_7", "92.000", "93.500", "269.000", "442.000", "25.000", "29.500"]
    audit = {
        "source": str(SOURCE),
        "output": str(OUT),
        "before": before,
        "after": {
            "paragraphs": len(checked.paragraphs),
            "tables": len(checked.tables),
            "text_chars": len(full_text),
        },
        "checks": {needle: full_text.count(needle) for needle in needles},
        "created_at": "2026-07-31",
        "note": "Existing v2.2 FINAL preserved. v2.3 Arbeitsstand appends a controlled reviewer/coach-readiness section.",
    }
    QA.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(audit, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
