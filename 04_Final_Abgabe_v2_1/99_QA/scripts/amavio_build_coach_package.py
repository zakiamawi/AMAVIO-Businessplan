from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    KeepTogether,
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


NAVY = "071523"
INK = "10243E"
BRASS = "D6B46A"
COPPER = "B97855"
CREAM = "F6EFE2"
PLATINUM = "D8D4C8"
MUTED = "514A41"


SOURCE_DOCS = [
    ("Coach Story", "AMAVIO_5_Minuten_Coach_Story_20260731.md"),
    ("P1 Gate Status", "AMAVIO_P1_Gate_Status_20260731.md"),
    ("Finanzmodell Bridge", "AMAVIO_Finanzmodell_Coach_Bridge_20260731.md"),
    ("Erstmandat & Pipeline", "AMAVIO_Ideales_Erstmandat_und_Pipeline_20260731.md"),
    ("Source Refresh", "AMAVIO_Source_Refresh_Markt_Recht_20260731.md"),
    ("14-Tage-Arbeitsplan", "AMAVIO_14_Tage_Arbeitsplan_20260731.md"),
    ("Submission Manifest", "AMAVIO_Coach_Submission_Manifest_20260731.md"),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, color: str = INK) -> None:
    cell.text = ""
    para = cell.paragraphs[0]
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def iter_blocks(markdown: str):
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("|") and "|" in line[1:]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = []
            for raw in table_lines:
                cells = [c.strip() for c in raw.strip("|").split("|")]
                if all(re.fullmatch(r":?-{3,}:?", c or "") for c in cells):
                    continue
                rows.append(cells)
            if rows:
                yield ("table", rows)
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            yield ("heading", (level, line[level:].strip()))
            i += 1
            continue
        if line.startswith("- "):
            items = []
            while i < len(lines) and lines[i].strip().startswith("- "):
                items.append(lines[i].strip()[2:].strip())
                i += 1
            yield ("bullets", items)
            continue
        if re.match(r"^\d+\.\s+", line):
            items = []
            while i < len(lines) and re.match(r"^\d+\.\s+", lines[i].strip()):
                items.append(re.sub(r"^\d+\.\s+", "", lines[i].strip()))
                i += 1
            yield ("numbers", items)
            continue
        para = [line]
        i += 1
        while i < len(lines):
            next_line = lines[i].rstrip()
            if not next_line or next_line.startswith("#") or next_line.startswith("|") or next_line.startswith("- ") or re.match(r"^\d+\.\s+", next_line.strip()):
                break
            para.append(next_line)
            i += 1
        yield ("para", " ".join(para).strip())


def clean_text(text: str) -> str:
    text = text.replace("`", "")
    text = text.replace("–", "-").replace("—", "-").replace("→", "->")
    text = text.replace("“", '"').replace("”", '"').replace("„", '"')
    return text


def add_docx_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(MUTED)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, color, before, after in [
        ("Heading 1", 16, INK, 14, 7),
        ("Heading 2", 13, INK, 10, 5),
        ("Heading 3", 11.5, COPPER, 8, 3),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_docx_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run("AMAVIO LIVING")
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string(BRASS)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Coach Readiness Paket")
    r.bold = True
    r.font.size = Pt(30)
    r.font.color.rgb = RGBColor.from_string(NAVY)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(22)
    sr = sub.add_run("Businessplan v2.2 | Finanzmodell v1.8 | Deck- und Pitch-Brücke")
    sr.font.size = Pt(12)
    sr.font.color.rgb = RGBColor.from_string(MUTED)

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows = [
        ("Stand", "31. Juli 2026"),
        ("Ziel", "Coach-fähige Entscheidungs- und Prüfunterlage"),
        ("Führende Zahlen", "Umsatz 93.500 / 269.000 / 442.000 EUR"),
        ("Wichtig", "Keine Website-Finalisierung in diesem Paket"),
    ]
    for row, (label, value) in zip(table.rows, rows):
        set_cell_text(row.cells[0], label, bold=True, color=INK)
        set_cell_text(row.cells[1], value, color=MUTED)
        set_cell_shading(row.cells[0], CREAM)
    doc.add_page_break()


def add_markdown_to_docx(doc: Document, title: str, markdown: str) -> None:
    doc.add_heading(title, level=1)
    for kind, value in iter_blocks(markdown):
        if kind == "heading":
            level, text = value
            if level == 1:
                continue
            doc.add_heading(clean_text(text), level=min(level, 3))
        elif kind == "para":
            if value.startswith("Quelle:") or value.startswith("Zweck:"):
                p = doc.add_paragraph(clean_text(value))
                p.runs[0].italic = True
            else:
                doc.add_paragraph(clean_text(value))
        elif kind == "bullets":
            for item in value:
                doc.add_paragraph(clean_text(item), style="List Bullet")
        elif kind == "numbers":
            for item in value:
                doc.add_paragraph(clean_text(item), style="List Number")
        elif kind == "table":
            rows = value
            max_cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=max_cols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r_idx, row in enumerate(rows):
                for c_idx in range(max_cols):
                    text = clean_text(row[c_idx]) if c_idx < len(row) else ""
                    set_cell_text(table.rows[r_idx].cells[c_idx], text, bold=(r_idx == 0), color=(INK if r_idx == 0 else MUTED))
                    if r_idx == 0:
                        set_cell_shading(table.rows[r_idx].cells[c_idx], CREAM)
            doc.add_paragraph()


def build_docx(source_dir: Path, out_path: Path) -> None:
    doc = Document()
    add_docx_styles(doc)
    add_docx_cover(doc)
    for index, (title, filename) in enumerate(SOURCE_DOCS):
        if index:
            doc.add_section(WD_SECTION_START.NEW_PAGE)
        add_markdown_to_docx(doc, title, (source_dir / filename).read_text(encoding="utf-8"))

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("AMAVIO Living - Coach Readiness Paket - 31. Juli 2026")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    if out_path.exists():
        raise FileExistsError(out_path)
    doc.save(out_path)


def pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("CoverBrand", alignment=TA_CENTER, fontSize=13, leading=15, textColor=colors.HexColor("#" + BRASS), spaceAfter=18))
    styles.add(ParagraphStyle("CoverTitle", alignment=TA_CENTER, fontSize=30, leading=34, textColor=colors.HexColor("#" + NAVY), spaceAfter=10))
    styles.add(ParagraphStyle("CoverSub", alignment=TA_CENTER, fontSize=11, leading=15, textColor=colors.HexColor("#" + MUTED), spaceAfter=24))
    styles.add(ParagraphStyle("H1A", fontSize=16, leading=20, textColor=colors.HexColor("#" + INK), spaceBefore=16, spaceAfter=8))
    styles.add(ParagraphStyle("H2A", fontSize=13, leading=16, textColor=colors.HexColor("#" + INK), spaceBefore=12, spaceAfter=6))
    styles.add(ParagraphStyle("H3A", fontSize=11, leading=14, textColor=colors.HexColor("#" + COPPER), spaceBefore=8, spaceAfter=4))
    styles.add(ParagraphStyle("BodyA", fontSize=9.2, leading=12.2, textColor=colors.HexColor("#" + MUTED), spaceAfter=5))
    return styles


def add_markdown_to_story(story: list, styles, title: str, markdown: str) -> None:
    story.append(Paragraph(title, styles["H1A"]))
    for kind, value in iter_blocks(markdown):
        if kind == "heading":
            level, text = value
            if level == 1:
                continue
            story.append(Paragraph(clean_text(text), styles["H2A" if level == 2 else "H3A"]))
        elif kind == "para":
            story.append(Paragraph(clean_text(value), styles["BodyA"]))
        elif kind in ("bullets", "numbers"):
            bullet_type = "1" if kind == "numbers" else "bullet"
            items = [ListItem(Paragraph(clean_text(item), styles["BodyA"])) for item in value]
            story.append(ListFlowable(items, bulletType=bullet_type, leftIndent=16))
        elif kind == "table":
            data = [[Paragraph(clean_text(cell), styles["BodyA"]) for cell in row] for row in value]
            tbl = Table(data, repeatRows=1, hAlign="LEFT")
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + CREAM)),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#" + INK)),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#" + PLATINUM)),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(tbl)
            story.append(Spacer(1, 6))


def build_pdf(source_dir: Path, out_path: Path) -> None:
    if out_path.exists():
        raise FileExistsError(out_path)
    styles = pdf_styles()
    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=letter,
        rightMargin=0.7 * inch,
        leftMargin=0.7 * inch,
        topMargin=0.7 * inch,
        bottomMargin=0.6 * inch,
    )
    story = [
        Spacer(1, 1.1 * inch),
        Paragraph("AMAVIO LIVING", styles["CoverBrand"]),
        Paragraph("Coach Readiness Paket", styles["CoverTitle"]),
        Paragraph("Businessplan v2.2 | Finanzmodell v1.8 | Deck- und Pitch-Brücke", styles["CoverSub"]),
        Table(
            [
                ["Stand", "31. Juli 2026"],
                ["Ziel", "Coach-fähige Entscheidungs- und Prüfunterlage"],
                ["Führende Zahlen", "Umsatz 93.500 / 269.000 / 442.000 EUR"],
                ["Wichtig", "Keine Website-Finalisierung in diesem Paket"],
            ],
            colWidths=[1.55 * inch, 4.8 * inch],
            style=TableStyle([
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#" + CREAM)),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#" + PLATINUM)),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#" + MUTED)),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("PADDING", (0, 0), (-1, -1), 7),
            ]),
        ),
        PageBreak(),
    ]
    for index, (title, filename) in enumerate(SOURCE_DOCS):
        if index:
            story.append(PageBreak())
        add_markdown_to_story(story, styles, title, (source_dir / filename).read_text(encoding="utf-8"))

    def footer(canvas, pdf_doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#" + MUTED))
        canvas.drawCentredString(letter[0] / 2, 0.35 * inch, f"AMAVIO Living - Coach Readiness Paket - Seite {pdf_doc.page}")
        canvas.restoreState()

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def audit_pdf(path: Path) -> dict:
    reader = PdfReader(str(path))
    chars = 0
    empty = []
    for idx, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        chars += len(text)
        if not text.strip():
            empty.append(idx)
    return {"page_count": len(reader.pages), "text_chars": chars, "empty_text_pages": empty}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-dir", type=Path, required=True)
    parser.add_argument("--out-dir", type=Path, required=True)
    args = parser.parse_args()

    docx_path = args.out_dir / "AMAVIO_Coach_Readiness_Paket_20260731.docx"
    pdf_path = args.out_dir / "AMAVIO_Coach_Readiness_Paket_20260731.pdf"
    build_docx(args.source_dir, docx_path)
    build_pdf(args.source_dir, pdf_path)
    result = {"docx": str(docx_path), "pdf": str(pdf_path), "pdf_audit": audit_pdf(pdf_path)}
    (args.out_dir / "AMAVIO_Coach_Readiness_Paket_20260731_QA.json").write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(result, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
