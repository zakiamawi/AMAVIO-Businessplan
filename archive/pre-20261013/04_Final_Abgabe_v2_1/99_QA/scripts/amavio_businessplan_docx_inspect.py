from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document


def main() -> int:
    if len(sys.argv) != 2:
        print("Usage: amavio_businessplan_docx_inspect.py <docx>", file=sys.stderr)
        return 2

    path = Path(sys.argv[1])
    doc = Document(path)

    headings: list[dict[str, object]] = []
    non_empty = 0
    for idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text:
            non_empty += 1
        style = para.style.name if para.style is not None else ""
        is_heading = (
            "Heading" in style
            or "KAPITEL" in text
            or text in {"Executive Summary", "Dokumentprofil und Leselogik"}
            or text.endswith("planung")
        )
        if text and is_heading:
            headings.append(
                {
                    "paragraph_index": idx,
                    "style": style,
                    "text": text[:180],
                }
            )

    full_text = "\n".join(p.text for p in doc.paragraphs)
    needles = ["v1.7", "v1_7", "92.000", "93.500", "269.000", "442.000", "25.000", "29.500"]
    checks = {needle: full_text.count(needle) for needle in needles}

    result = {
        "path": str(path),
        "paragraphs": len(doc.paragraphs),
        "non_empty_paragraphs": non_empty,
        "tables": len(doc.tables),
        "headings_found": len(headings),
        "headings_preview": headings[:120],
        "checks": checks,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
