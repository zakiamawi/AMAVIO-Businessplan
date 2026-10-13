from __future__ import annotations

import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image as RLImage
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path("/Users/zakia/Desktop/Businessplan AKTUELL")
OUT = ROOT / "08_Website/brand/AMAVIO_One_Brandbook_Master_V2_Designveredelung_20260930"
ASSETS = OUT / "assets"
RENDER = OUT / "rendered_pages"
MASTER_MD = OUT / "AMAVIO_One_Brandbook_Master_Arbeitsfassung_20260930.md"
DOCX = OUT / "AMAVIO_One_Brandbook_Master_Arbeitsfassung_20260930.docx"
PDF = OUT / "AMAVIO_One_Brandbook_Master_Arbeitsfassung_20260930.pdf"
QUICK_PDF = OUT / "AMAVIO_One_Brand_Quick_Reference_20260930.pdf"

NAVY = "0B1F3A"
STAGE = "081527"
CREAM = "F7F4EC"
WHITE = "FFFFFF"
INK = "222222"
GRAY = "6F7680"
GOLD = "C6A35A"
GOLD_DEEP = "9C7D3F"
BRONZE = "7D6130"
STONE = "D9D0C3"
LINE = "E2DCC9"
CONFIRMED = "234C78"
ATTENTION = "A94F45"


def rgb(hex_value: str) -> RGBColor:
    value = hex_value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill.lstrip("#"))


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **edges) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, data in edges.items():
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        for key, value in data.items():
            el.set(qn(f"w:{key}"), str(value))


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_paragraph_border(paragraph, *, top=None, bottom=None, start=None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    for edge, data in (("top", top), ("bottom", bottom), ("start", start)):
        if not data:
            continue
        el = p_bdr.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            p_bdr.append(el)
        for key, value in data.items():
            el.set(qn(f"w:{key}"), str(value))


def set_keep_with_next(paragraph, value: bool = True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepNext"))
    if keep is None:
        keep = OxmlElement("w:keepNext")
        p_pr.append(keep)
    keep.set(qn("w:val"), "1" if value else "0")


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_inline(paragraph, text: str, *, color: str | None = None, size: float | None = None) -> None:
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\([^)]+\))")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            if color:
                run.font.color.rgb = rgb(color)
            if size:
                run.font.size = Pt(size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Arial"
            run.font.size = Pt((size or 10) - 0.4)
            run.font.color.rgb = rgb(BRONZE)
        else:
            label = token[1 : token.index("]")]
            target = token[token.index("(") + 1 : -1]
            run = paragraph.add_run(label)
            run.underline = True
            run.font.color.rgb = rgb(BRONZE)
            run._r.set(qn("w:rsidRPr"), "00000000")
        if color and not token.startswith("`"):
            run.font.color.rgb = rgb(color)
        if size and not token.startswith("`"):
            run.font.size = Pt(size)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        if color:
            run.font.color.rgb = rgb(color)
        if size:
            run.font.size = Pt(size)


def configure_section(section) -> None:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.95)
    section.right_margin = Cm(1.65)
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.7)


def configure_document(doc: Document) -> None:
    configure_section(doc.sections[0])
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(4.5)
    normal.paragraph_format.line_spacing = 1.18

    h1 = styles["Heading 1"]
    h1.font.name = "Georgia"
    h1.font.size = Pt(21)
    h1.font.bold = False
    h1.font.color.rgb = rgb(NAVY)
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(14)
    h1.paragraph_format.keep_with_next = True

    h2 = styles["Heading 2"]
    h2.font.name = "Arial"
    h2.font.size = Pt(12.5)
    h2.font.bold = True
    h2.font.color.rgb = rgb(NAVY)
    h2.paragraph_format.space_before = Pt(11)
    h2.paragraph_format.space_after = Pt(5)
    h2.paragraph_format.keep_with_next = True

    h3 = styles["Heading 3"]
    h3.font.name = "Arial"
    h3.font.size = Pt(10.5)
    h3.font.bold = True
    h3.font.color.rgb = rgb(BRONZE)
    h3.paragraph_format.space_before = Pt(9)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(9.5)
        style.font.color.rgb = rgb(INK)
        style.paragraph_format.left_indent = Cm(0.55)
        style.paragraph_format.first_line_indent = Cm(-0.25)
        style.paragraph_format.space_after = Pt(3)

    if "AMAVIO Lead" not in styles:
        lead = styles.add_style("AMAVIO Lead", WD_STYLE_TYPE.PARAGRAPH)
    else:
        lead = styles["AMAVIO Lead"]
    lead.font.name = "Georgia"
    lead.font.size = Pt(11.2)
    lead.font.color.rgb = rgb(NAVY)
    lead.paragraph_format.space_after = Pt(10)
    lead.paragraph_format.line_spacing = 1.26

    props = doc.core_properties
    props.title = "AMAVIO One Brandbook Master V2 Designveredelung 2026-09-30"
    props.subject = "AMAVIO One Design-, Sprach- und Kommunikationssystem"
    props.author = "AMAVIO One"
    props.keywords = "AMAVIO One, Brandbook, Design System, Working Draft"
    props.comments = "V2 Designveredelung auf inhaltsgleichem Master; interne Arbeitsfassung; externe Freigabe und Founder Gates offen."


def add_header_footer(section) -> None:
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.25), WD_TAB_ALIGNMENT.RIGHT)
    set_paragraph_border(p, bottom={"val": "single", "sz": 5, "space": 4, "color": LINE})
    r = p.add_run("AMAVIO ONE")
    r.font.name = "Arial"
    r.font.size = Pt(7.2)
    r.font.bold = True
    r.font.color.rgb = rgb(BRONZE)
    r = p.add_run("\tBRAND SYSTEM · MASTER V2")
    r.font.name = "Arial"
    r.font.size = Pt(7.2)
    r.font.bold = True
    r.font.color.rgb = rgb(GRAY)

    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.25), WD_TAB_ALIGNMENT.RIGHT)
    set_paragraph_border(p, top={"val": "single", "sz": 5, "space": 4, "color": LINE})
    r = p.add_run("MASTER-ARBEITSFASSUNG · DESIGNVEREDELUNG · 30.09.2026")
    r.font.name = "Arial"
    r.font.size = Pt(7.0)
    r.font.color.rgb = rgb(GRAY)
    p.add_run("\t")
    add_page_field(p)
    for run in p.runs:
        run.font.name = "Arial"
        run.font.size = Pt(7.0)
        run.font.color.rgb = rgb(GRAY)


def add_cover(doc: Document) -> None:
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    panel = doc.add_table(rows=1, cols=1)
    panel.alignment = WD_TABLE_ALIGNMENT.CENTER
    panel.autofit = False
    row = panel.rows[0]
    row.height = Cm(25.45)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    cell = row.cells[0]
    cell.width = Cm(17.35)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    shade(cell, STAGE)
    set_cell_margins(cell, top=430, start=460, bottom=360, end=460)
    set_cell_border(
        cell,
        top={"val": "nil"}, bottom={"val": "nil"},
        start={"val": "nil"}, end={"val": "nil"},
        insideH={"val": "nil"}, insideV={"val": "nil"},
    )

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(38)
    logo = ASSETS / "logo-light-on-navy.png"
    if logo.exists():
        p.add_run().add_picture(str(logo), width=Inches(3.05))
    else:
        r = p.add_run("AMAVIO ONE")
        r.font.name = "Georgia"
        r.font.size = Pt(28)
        r.font.color.rgb = rgb(CREAM)

    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("BRAND SYSTEM · MASTER V2 · 2026")
    r.font.name = "Arial"
    r.font.size = Pt(8.2)
    r.font.bold = True
    r.font.color.rgb = rgb(GOLD)

    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    set_paragraph_border(p, start={"val": "single", "sz": 18, "space": 10, "color": GOLD})
    r = p.add_run("Design.\nSprache.\nSystem.")
    r.font.name = "Georgia"
    r.font.size = Pt(31)
    r.font.color.rgb = rgb(CREAM)

    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(26)
    r = p.add_run("Ein konsistentes Marken- und Anwendungssystem für arbeitgeberseitiges Arrival Management.")
    r.font.name = "Arial"
    r.font.size = Pt(11.2)
    r.font.color.rgb = rgb("DDE4EA")

    meta = cell.add_table(rows=1, cols=2)
    meta.autofit = False
    meta.columns[0].width = Cm(7.2)
    meta.columns[1].width = Cm(7.2)
    meta_values = [
        ("MARKE", "AMAVIO One"),
        ("STATUS", "Master-Arbeitsfassung"),
    ]
    for c, (label, value) in zip(meta.rows[0].cells, meta_values):
        shade(c, NAVY)
        set_cell_border(c, top={"val": "single", "sz": 8, "color": GOLD})
        set_cell_margins(c, top=150, start=0, bottom=150, end=180)
        p = c.paragraphs[0]
        r = p.add_run(label + "\n")
        r.font.name = "Arial"
        r.font.size = Pt(6.8)
        r.font.bold = True
        r.font.color.rgb = rgb(GOLD)
        r = p.add_run(value)
        r.font.name = "Arial"
        r.font.size = Pt(9.2)
        r.font.bold = True
        r.font.color.rgb = rgb(CREAM)

    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(34)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("ANKUNFT STEUERBAR MACHEN.")
    r.font.name = "Arial"
    r.font.size = Pt(8.5)
    r.font.bold = True
    r.font.color.rgb = rgb(GOLD)

    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Intern · Stand 30.09.2026 · nicht extern freigegeben")
    r.font.name = "Arial"
    r.font.size = Pt(7.6)
    r.font.color.rgb = rgb("AEB9C6")
    doc.add_page_break()


def add_contents(doc: Document, lines: list[str]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("SYSTEMÜBERSICHT")
    r.font.name = "Arial"
    r.font.size = Pt(7.5)
    r.font.bold = True
    r.font.color.rgb = rgb(BRONZE)
    p = doc.add_paragraph(style="Heading 1")
    add_inline(p, "Inhalt und Leselogik")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(13)
    add_inline(p, "Vier Ebenen führen von der Markenidee bis zur kontrollierten Anwendung: Marke → Ausdruck → Anwendung → Governance.", color=GRAY, size=9.6)

    entries = []
    for line in lines:
        if line.startswith("## "):
            title = line[3:].strip()
            match = re.match(r"^(\d+)\.\s*(.+)$", title)
            entries.append((int(match.group(1)), match.group(2)) if match else (0, title))

    groups = [(0, 0, "ORIENTIERUNG"), (1, 5, "MARKE"), (6, 12, "AUSDRUCK"), (13, 18, "ANWENDUNG"), (19, 22, "GOVERNANCE")]
    def group_for(number: int) -> str:
        return next(label for start, end, label in groups if start <= number <= end)

    grid = doc.add_table(rows=1, cols=2)
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid.autofit = False
    halves = [entries[:12], entries[12:]]
    for col, items in enumerate(halves):
        cell = grid.cell(0, col)
        cell.width = Cm(8.35)
        set_cell_margins(cell, top=0, start=0 if col == 0 else 180, bottom=0, end=180 if col == 0 else 0)
        set_cell_border(cell, top={"val": "nil"}, bottom={"val": "nil"}, start={"val": "nil"}, end={"val": "nil"})
        cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)
        previous_group = None
        for number, title in items:
            group = group_for(number)
            if group != previous_group:
                p = cell.add_paragraph()
                p.paragraph_format.space_before = Pt(6 if previous_group else 0)
                p.paragraph_format.space_after = Pt(3)
                r = p.add_run(group)
                r.font.name = "Arial"
                r.font.size = Pt(6.7)
                r.font.bold = True
                r.font.color.rgb = rgb(BRONZE)
                previous_group = group
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            set_paragraph_border(p, bottom={"val": "single", "sz": 4, "space": 4, "color": LINE})
            r = p.add_run(f"{number:02d}  ")
            r.font.name = "Arial"
            r.font.size = Pt(8.4)
            r.font.bold = True
            r.font.color.rgb = rgb(BRONZE)
            r = p.add_run(title)
            r.font.name = "Arial"
            r.font.size = Pt(8.8)
            r.font.color.rgb = rgb(NAVY)
    doc.add_page_break()


def chapter_group(number: int) -> str:
    if number == 0:
        return "ORIENTIERUNG"
    if number <= 5:
        return "MARKENFUNDAMENT"
    if number <= 12:
        return "VISUELLES SYSTEM"
    if number <= 18:
        return "ANWENDUNG"
    return "GOVERNANCE"


def add_chapter_heading(doc: Document, title: str, fallback_number: int, start_new_page: bool = False) -> None:
    if start_new_page:
        marker = doc.add_paragraph()
        marker.paragraph_format.page_break_before = True
        marker.paragraph_format.space_before = Pt(0)
        marker.paragraph_format.space_after = Pt(0)
        marker.paragraph_format.line_spacing = Pt(1)
        marker.add_run("").font.size = Pt(1)
    match = re.match(r"^(\d+)\.\s*(.+)$", title)
    number = int(match.group(1)) if match else 0
    clean_title = match.group(2) if match else title
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    left, right = table.rows[0].cells
    left.width = Cm(14.5)
    right.width = Cm(2.6)
    for cell in (left, right):
        set_cell_border(cell, top={"val": "nil"}, start={"val": "nil"}, end={"val": "nil"}, bottom={"val": "single", "sz": 10, "color": GOLD})
        set_cell_margins(cell, top=0, start=0, bottom=110, end=0)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM

    p = left.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(chapter_group(number))
    r.font.name = "Arial"
    r.font.size = Pt(7.1)
    r.font.bold = True
    r.font.color.rgb = rgb(BRONZE)
    p = left.add_paragraph()
    p.style = doc.styles["Heading 1"]
    p.paragraph_format.space_after = Pt(0)
    add_inline(p, clean_title)

    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{number:02d}")
    r.font.name = "Georgia"
    r.font.size = Pt(30)
    r.font.color.rgb = rgb(GOLD_DEEP)


def add_logo_panel(doc: Document) -> None:
    doc.add_paragraph("Aktuelle AMAVIO-One-Logoassets", style="Heading 2")
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    items = [
        ("PRIMARY · helle Fläche", ASSETS / "logo-primary.png"),
        ("LIGHT · dunkle Fläche", ASSETS / "logo-light-on-navy.png"),
        ("DARK · eingebettete Fläche", ASSETS / "logo-dark.png"),
        ("SIGNET · kleiner Kontext", ASSETS / "signet.png"),
    ]
    for index, (cell, (label, image_path)) in enumerate(zip([c for r in table.rows for c in r.cells], items)):
        shade(cell, "FAF8F2" if index % 3 else "F2EEE4")
        set_cell_margins(cell, top=150, start=150, bottom=120, end=150)
        set_cell_border(cell, top={"val": "single", "sz": 6, "color": GOLD}, bottom={"val": "nil"}, start={"val": "nil"}, end={"val": "nil"})
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if image_path.exists():
            width = Inches(2.55) if "signet" not in image_path.name else Inches(0.85)
            p.add_run().add_picture(str(image_path), width=width)
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        r.font.name = "Arial"
        r.font.size = Pt(7.5)
        r.font.bold = True
        r.font.color.rgb = rgb(BRONZE)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        values = [v.strip().replace("\\|", "|") for v in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", v) for v in values):
            rows.append(values)
        i += 1
    return rows, i


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=1, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header = table.rows[0]
    set_row_cant_split(header)
    for idx in range(cols):
        cell = header.cells[idx]
        cell.text = rows[0][idx] if idx < len(rows[0]) else ""
        shade(cell, STAGE)
        set_cell_margins(cell, top=90, start=110, bottom=90, end=110)
        set_cell_border(cell, top={"val": "nil"}, start={"val": "nil"}, end={"val": "nil"}, bottom={"val": "single", "sz": 10, "color": GOLD})
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                run.font.name = "Arial"
                run.font.size = Pt(8.2)
                run.font.bold = True
                run.font.color.rgb = rgb(CREAM)
    set_repeat_table_header(header)

    hex_col = None
    for idx, label in enumerate(rows[0]):
        if label.strip().upper() == "HEX":
            hex_col = idx
    for row_index, values in enumerate(rows[1:]):
        row = table.add_row()
        set_row_cant_split(row)
        cells = row.cells
        for idx in range(cols):
            text = values[idx] if idx < len(values) else ""
            cells[idx].text = ""
            set_cell_margins(cells[idx], top=76, start=110, bottom=76, end=110)
            set_cell_border(cells[idx], top={"val": "nil"}, start={"val": "nil"}, end={"val": "nil"}, bottom={"val": "single", "sz": 4, "color": LINE})
            if row_index % 2:
                shade(cells[idx], "F7F4EC")
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, text, color=INK, size=8.2)
            if idx == 0 and p.runs:
                p.runs[0].bold = True
                p.runs[0].font.color.rgb = rgb(NAVY)
            if hex_col is not None and idx == hex_col:
                match = re.search(r"#[0-9A-Fa-f]{6}", text)
                if match:
                    fill = match.group(0).lstrip("#")
                    shade(cells[idx], fill)
                    lum = sum(int(fill[i : i + 2], 16) * weight for i, weight in zip((0, 2, 4), (0.2126, 0.7152, 0.0722)))
                    color = WHITE if lum < 125 else INK
                    for run in p.runs:
                        run.font.color.rgb = rgb(color)
                        run.font.bold = True
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    table.rows[-1].cells[0].paragraphs[-1].paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_body_from_markdown(doc: Document, lines: list[str]) -> None:
    i = 0
    first_h2 = True
    chapter_index = 0
    chapter_lead_pending = False
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            i += 1
            continue
        if line.startswith("## "):
            title = line[3:].strip()
            chapter_index += 1
            number_match = re.match(r"^(\d+)\.\s*", title)
            chapter_number = int(number_match.group(1)) if number_match else 0
            editorial_flow_chapters = {10, 12, 13, 15, 17, 18}
            add_chapter_heading(
                doc,
                title,
                chapter_index,
                start_new_page=not first_h2 and chapter_number not in editorial_flow_chapters,
            )
            first_h2 = False
            chapter_lead_pending = True
            if title.startswith("6. Logo-System"):
                add_logo_panel(doc)
                chapter_lead_pending = False
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            set_paragraph_border(p, top={"val": "single", "sz": 5, "space": 8, "color": LINE})
            add_inline(p, line[4:].strip())
            chapter_lead_pending = False
            i += 1
            continue
        if line.startswith("#### "):
            p = doc.add_paragraph(style="Heading 3")
            add_inline(p, line[5:].strip())
            chapter_lead_pending = False
            i += 1
            continue
        if line.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            chapter_lead_pending = False
            continue
        if line == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(5)
            set_paragraph_border(p, bottom={"val": "single", "sz": 8, "space": 5, "color": GOLD})
            i += 1
            continue
        if line.startswith("> "):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].strip().lstrip(">").strip())
                i += 1
            table = doc.add_table(rows=1, cols=1)
            cell = table.cell(0, 0)
            shade(cell, NAVY)
            set_cell_border(cell, start={"val": "single", "sz": 18, "color": GOLD})
            set_cell_margins(cell, top=170, start=210, bottom=170, end=190)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, " ".join(quote_lines), color=CREAM, size=11.2)
            for run in p.runs:
                run.font.name = "Georgia"
                run.italic = True
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
            chapter_lead_pending = False
            continue
        if line.startswith("- "):
            text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.55)
            p.paragraph_format.first_line_indent = Cm(-0.32)
            p.paragraph_format.space_after = Pt(2.4)
            r = p.add_run("—  ")
            r.font.name = "Arial"
            r.font.bold = True
            r.font.color.rgb = rgb(GOLD_DEEP)
            add_inline(p, text)
            chapter_lead_pending = False
            i += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            number_match = re.match(r"^(\d+)\.\s+", line)
            text = re.sub(r"^\d+\.\s+", "", line)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.62)
            p.paragraph_format.first_line_indent = Cm(-0.42)
            p.paragraph_format.space_after = Pt(2.4)
            r = p.add_run(f"{number_match.group(1).zfill(2)}  ")
            r.font.name = "Arial"
            r.font.bold = True
            r.font.color.rgb = rgb(BRONZE)
            add_inline(p, text)
            chapter_lead_pending = False
            i += 1
            continue
        if re.match(r"^\*\*[^*]+:\*\*", line) and len(line) < 160:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(5)
            add_inline(p, line)
            chapter_lead_pending = False
            i += 1
            continue

        para_lines = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt:
                i += 1
                break
            if nxt.startswith(("#", "|", ">", "- ")) or nxt == "---" or re.match(r"^\d+\.\s+", nxt):
                break
            para_lines.append(nxt)
            i += 1
        p = doc.add_paragraph(style="AMAVIO Lead" if chapter_lead_pending else None)
        add_inline(p, " ".join(para_lines), size=11.2 if chapter_lead_pending else None)
        chapter_lead_pending = False


def build_docx() -> None:
    lines = MASTER_MD.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_contents(doc, lines)
    add_body_from_markdown(doc, lines)
    for section in doc.sections:
        configure_section(section)
        add_header_footer(section)
    doc.save(DOCX)


def build_quick_reference_pdf() -> None:
    base = getSampleStyleSheet()
    styles = {
        "eyebrow": ParagraphStyle("eyebrow", base["Normal"], fontName="Helvetica-Bold", fontSize=7.2, leading=9, textColor=colors.HexColor("#7D6130"), spaceAfter=4),
        "title": ParagraphStyle("title", base["Title"], fontName="Times-Roman", fontSize=23, leading=24, textColor=colors.HexColor("#0B1F3A"), spaceAfter=6, alignment=TA_LEFT),
        "h": ParagraphStyle("h", base["Heading2"], fontName="Helvetica-Bold", fontSize=8.8, leading=10.5, textColor=colors.HexColor("#0B1F3A"), spaceBefore=4, spaceAfter=3),
        "body": ParagraphStyle("body", base["BodyText"], fontName="Helvetica", fontSize=7.4, leading=9.6, textColor=colors.HexColor("#222222"), spaceAfter=2.5),
        "small": ParagraphStyle("small", base["BodyText"], fontName="Helvetica", fontSize=6.6, leading=8.2, textColor=colors.HexColor("#6F7680"), spaceAfter=2),
    }

    def header_footer(canvas, document):
        canvas.saveState()
        w, h = A4
        canvas.setFillColor(colors.HexColor("#081527"))
        canvas.rect(0, h - 12 * mm, w, 12 * mm, fill=1, stroke=0)
        canvas.setFillColor(colors.HexColor("#C6A35A"))
        canvas.setFont("Helvetica-Bold", 7)
        canvas.drawString(16 * mm, h - 7.4 * mm, "AMAVIO ONE  ·  BRAND QUICK REFERENCE")
        canvas.setFillColor(colors.HexColor("#6F7680"))
        canvas.setFont("Helvetica", 6.5)
        canvas.drawString(16 * mm, 8 * mm, "Master-Arbeitsfassung · 30.09.2026 · intern, nicht extern freigegeben")
        canvas.drawRightString(w - 16 * mm, 8 * mm, "1 / 1")
        canvas.restoreState()

    doc = SimpleDocTemplate(
        str(QUICK_PDF),
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=18 * mm,
        bottomMargin=13 * mm,
        title="AMAVIO One Brand Quick Reference 20260930",
        author="AMAVIO One",
    )
    story = []
    story.append(Paragraph("MASTER-ARBEITSFASSUNG · EINE SEITE", styles["eyebrow"]))
    story.append(Paragraph("AMAVIO One Brand System", styles["title"]))
    story.append(Paragraph("Servicegeführtes, arbeitgeberseitiges Arrival Management nach der Personalzusage. ACC und App sind getrennte Produktbausteine; aktuelle digitale Stände sind lokal/synthetisch, nicht produktiv.", styles["body"]))
    story.append(Spacer(1, 2 * mm))

    left = [
        Paragraph("NAMING & PRODUKT", styles["h"]),
        Paragraph("<b>AMAVIO One</b> ist Dachmarke. <b>Arrival Control Center</b>: Arbeitgeber/AMAVIO-Steuerung. <b>AMAVIO One App</b>: Fachkraft, Wohnanbieter, lokaler Partner. <b>AMAVIO Living</b> nur nach Founder-Freigabe als Modul.", styles["body"]),
        Paragraph("FARBE", styles["h"]),
    ]
    swatches = [
        ["Navy", "#0B1F3A"], ["Stage", "#081527"], ["Cream", "#F7F4EC"],
        ["Ink", "#222222"], ["Gold", "#C6A35A"], ["Bronze", "#7D6130"],
        ["Confirmed", "#234C78"], ["Attention", "#A94F45"],
    ]
    swatch_table = Table(swatches, colWidths=[22 * mm, 25 * mm], rowHeights=5.6 * mm)
    commands = [("FONT", (0, 0), (-1, -1), "Helvetica", 6.8), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("BOX", (0, 0), (-1, -1), 0.3, colors.HexColor("#D9D0C3")), ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D9D0C3"))]
    for idx, (_, hx) in enumerate(swatches):
        commands.append(("BACKGROUND", (1, idx), (1, idx), colors.HexColor(hx)))
        commands.append(("TEXTCOLOR", (1, idx), (1, idx), colors.white if hx in {"#0B1F3A", "#081527", "#222222", "#234C78", "#A94F45", "#7D6130"} else colors.HexColor("#222222")))
        commands.append(("FONTNAME", (1, idx), (1, idx), "Helvetica-Bold"))
    swatch_table.setStyle(TableStyle(commands))
    left.append(swatch_table)
    left += [
        Paragraph("TYPOGRAFIE", styles["h"]),
        Paragraph("Website: Geist + Fraunces, Mono nur für kurze Labels. App: Geist/Inter. Office: Arial; Georgia nur als kurzer Akzent. Keine kritische Information in Mini-Schrift.", styles["body"]),
        Paragraph("FORM", styles["h"]),
        Paragraph("Weißraum, feine 1-px-Linien, klare Hierarchie. Radien 0/4/8 px. Arrival Line, Gate, Decision Ledger und Partner Layer sind Struktur – keine Dekoration.", styles["body"]),
    ]

    right = [
        Paragraph("SPRACHE", styles["h"]),
        Paragraph("Bevorzugt: <b>steuerbar, sichtbar, prüfbar, Status, Zuständigkeit, Freigabe, Partnerpfad, definierter Scope</b>. Fakten, Plan, Demo und Nachweis offen trennen.", styles["body"]),
        Paragraph("NICHT BEHAUPTEN", styles["h"]),
        Paragraph("Keine Personal-/Wohnungsvermittlung; keine fertige Plattform, App oder KI; keine Garantie für Wohnraum, Einreise, Anerkennung, Förderung, ROI, Wirkung oder Bindung.", styles["body"]),
        Paragraph("LOGO", styles["h"]),
        Paragraph("Nur vorhandene SVGs nutzen. Primary auf hell, Light auf dunkel. Nicht stauchen, neu färben, nachzeichnen, mit Effekten versehen oder das Signet kontextlos als Hauptlogo einsetzen.", styles["body"]),
        Paragraph("8-PUNKT-FINALCHECK", styles["h"]),
        Paragraph("1. AMAVIO One ist Dachmarke.<br/>2. Reifegrad sichtbar.<br/>3. ACC und App getrennt.<br/>4. Claims und Belege aktuell.<br/>5. Navy/Gold/Creme führt.<br/>6. Lesbarkeit, Kontrast, Weißraum stimmen.<br/>7. Datei, Textauszug und Render geprüft.<br/>8. Externe Aktion separat freigegeben.", styles["body"]),
        Paragraph("FOUNDER GATES", styles["h"]),
        Paragraph("AMAVIO Living als Modul · Rechtsträger/Markenrecht · öffentlicher Claim · Logo-Reinzeichnung · Executive Living · dauerhafte Digitalschriften.", styles["small"]),
    ]

    grid = Table([[left, right]], colWidths=[84 * mm, 84 * mm], hAlign="LEFT")
    grid.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (0, 0), 0),
        ("RIGHTPADDING", (0, 0), (0, 0), 7 * mm),
        ("LEFTPADDING", (1, 0), (1, 0), 7 * mm),
        ("RIGHTPADDING", (1, 0), (1, 0), 0),
        ("LINEBEFORE", (1, 0), (1, 0), 0.6, colors.HexColor("#C6A35A")),
    ]))
    story.append(grid)
    doc.build(story, onFirstPage=header_footer)


def main() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    RENDER.mkdir(parents=True, exist_ok=True)
    build_docx()
    build_quick_reference_pdf()
    print(DOCX)
    print(QUICK_PDF)


if __name__ == "__main__":
    main()
