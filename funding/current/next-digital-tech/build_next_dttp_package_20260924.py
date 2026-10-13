from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent

DEEP_INK = "101414"
DEEP_FOREST = "031212"
DARK_TEAL = "064E55"
OP_TEAL = "155B5D"
WARM_IVORY = "F8F1E8"
CLOUD_WHITE = "F7F8F4"
SOFT_STONE = "D9D0C3"
SOFT_STEEL = "DDE7E3"
CHAMPAGNE = "D7B98D"
MUTED_MINT = "4FD6B8"
CORAL = "F2765B"
WHITE = "FFFFFF"


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=SOFT_STONE, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_document(doc, landscape=False, footer_note=""):
    section = doc.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Cm(29.7)
        section.page_height = Cm(21.0)
        section.top_margin = Cm(1.25)
        section.bottom_margin = Cm(1.0)
        section.left_margin = Cm(1.25)
        section.right_margin = Cm(1.25)
    else:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.65)
        section.bottom_margin = Cm(1.55)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(9.6 if not landscape else 8.2)
    normal.font.color.rgb = RGBColor.from_string(DEEP_INK)
    normal.paragraph_format.space_after = Pt(4.5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in (
        ("Title", 25, DEEP_INK),
        ("Heading 1", 16, DARK_TEAL),
        ("Heading 2", 11.5, OP_TEAL),
        ("Heading 3", 9.5, DARK_TEAL),
    ):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(9 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("AMAVIO")
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(DARK_TEAL)
    r2 = p.add_run("  |  WORKFORCE ARRIVAL MANAGEMENT")
    r2.font.name = "Aptos"
    r2.font.size = Pt(7.2)
    r2.font.color.rgb = RGBColor.from_string(OP_TEAL)

    footer = section.footer
    footer.is_linked_to_previous = False
    table = footer.add_table(rows=1, cols=2, width=section.page_width - section.left_margin - section.right_margin)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_table_borders(table)
    left, right = table.rows[0].cells
    left.width = Cm(14 if not landscape else 22)
    right.width = Cm(3)
    lp = left.paragraphs[0]
    lr = lp.add_run(footer_note)
    lr.font.name = "Aptos"
    lr.font.size = Pt(6.8)
    lr.font.color.rgb = RGBColor.from_string(OP_TEAL)
    add_page_number(right.paragraphs[0])
    for run in right.paragraphs[0].runs:
        run.font.size = Pt(7)
        run.font.color.rgb = RGBColor.from_string(OP_TEAL)


def add_top_rule(doc, color=CHAMPAGNE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_shading(cell, color)
    set_cell_margins(cell, 0, 0, 0, 0)
    cell.height = Cm(0.08)
    remove_table_borders(table)


def strip_md(text):
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"`(.*?)`", r"\1", text)
    return text.strip()


def add_bold_markdown(paragraph, text, color=None, size=None):
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        clean = part[2:-2] if bold else part
        run = paragraph.add_run(clean)
        run.bold = bold
        run.font.name = "Aptos"
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if size:
            run.font.size = Pt(size)


def add_markdown_table(doc, rows):
    if not rows:
        return
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, value in enumerate(rows[0]):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, DARK_TEAL)
        set_cell_margins(cell, 90, 100, 90, 100)
        set_cell_border(cell, DARK_TEAL)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(strip_md(value))
        r.bold = True
        r.font.name = "Aptos"
        r.font.size = Pt(7.8)
        r.font.color.rgb = RGBColor.from_string(WHITE)
    for row_values in rows[1:]:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            cell = cells[i]
            set_cell_shading(cell, CLOUD_WHITE if len(table.rows) % 2 else WARM_IVORY)
            set_cell_margins(cell, 75, 90, 75, 90)
            set_cell_border(cell, SOFT_STONE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(strip_md(value))
            r.font.name = "Aptos"
            r.font.size = Pt(7.4)
            r.font.color.rgb = RGBColor.from_string(DEEP_INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def markdown_to_docx(source, target, footer_note, subtitle=None, compact=False):
    lines = Path(source).read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc, footer_note=footer_note)
    if compact:
        section = doc.sections[0]
        section.top_margin = Cm(1.35)
        section.bottom_margin = Cm(1.35)
        section.left_margin = Cm(1.75)
        section.right_margin = Cm(1.75)
        normal = doc.styles["Normal"]
        normal.font.size = Pt(9.1)
        normal.paragraph_format.space_after = Pt(3.2)
        normal.paragraph_format.line_spacing = 1.02
        for name, size in (("Title", 22), ("Heading 1", 14.5), ("Heading 2", 10.5)):
            style = doc.styles[name]
            style.font.size = Pt(size)
            style.paragraph_format.space_before = Pt(6 if name != "Title" else 0)
            style.paragraph_format.space_after = Pt(3)
    add_top_rule(doc)

    table_rows = []
    first_h1 = True
    i = 0
    while i < len(lines):
        raw = lines[i].rstrip()
        line = raw.strip()
        if not line:
            if table_rows:
                add_markdown_table(doc, table_rows)
                table_rows = []
            i += 1
            continue
        if line.startswith("|") and line.endswith("|"):
            values = [v.strip() for v in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", v or "") for v in values):
                i += 1
                continue
            table_rows.append(values)
            i += 1
            continue
        if table_rows:
            add_markdown_table(doc, table_rows)
            table_rows = []
        if line.startswith("# "):
            title = strip_md(line[2:])
            p = doc.add_paragraph(style="Title")
            p.paragraph_format.space_before = Pt(9 if first_h1 else 18)
            p.add_run(title)
            if first_h1 and subtitle:
                sp = doc.add_paragraph()
                sp.paragraph_format.space_after = Pt(12)
                sr = sp.add_run(subtitle)
                sr.font.name = "Aptos"
                sr.font.size = Pt(10)
                sr.font.color.rgb = RGBColor.from_string(OP_TEAL)
            first_h1 = False
        elif line.startswith("## "):
            doc.add_paragraph(strip_md(line[3:]), style="Heading 1")
        elif line.startswith("### "):
            doc.add_paragraph(strip_md(line[4:]), style="Heading 2")
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.45)
            p.paragraph_format.first_line_indent = Cm(-0.18)
            p.paragraph_format.space_after = Pt(2)
            add_bold_markdown(p, line[2:])
        elif re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Cm(0.45)
            p.paragraph_format.first_line_indent = Cm(-0.18)
            p.paragraph_format.space_after = Pt(2)
            add_bold_markdown(p, re.sub(r"^\d+\.\s", "", line))
        elif line.startswith("> "):
            table = doc.add_table(rows=1, cols=1)
            cell = table.cell(0, 0)
            set_cell_shading(cell, SOFT_STEEL)
            set_cell_margins(cell, 160, 180, 160, 180)
            set_cell_border(cell, MUTED_MINT, "10")
            p = cell.paragraphs[0]
            r = p.add_run(strip_md(line[2:]))
            r.bold = True
            r.font.name = "Aptos Display"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor.from_string(DARK_TEAL)
        elif line.startswith("**") and line.endswith("**"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(strip_md(line))
            r.bold = True
            r.font.color.rgb = RGBColor.from_string(OP_TEAL)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.widow_control = True
            add_bold_markdown(p, line)
        i += 1
    if table_rows:
        add_markdown_table(doc, table_rows)
    doc.save(target)


def add_bmc_cell(cell, title, bullets, accent=DARK_TEAL, font_size=7.2):
    set_cell_shading(cell, CLOUD_WHITE)
    set_cell_border(cell, SOFT_STONE, "7")
    set_cell_margins(cell, 95, 105, 80, 105)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title.upper())
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(accent)
    for bullet in bullets:
        p = cell.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.23)
        p.paragraph_format.first_line_indent = Cm(-0.15)
        p.paragraph_format.space_after = Pt(1.2)
        p.paragraph_format.line_spacing = 0.95
        r = p.add_run("• " + bullet)
        r.font.name = "Aptos"
        r.font.size = Pt(font_size)
        r.font.color.rgb = RGBColor.from_string(DEEP_INK)


def build_bmc(target):
    doc = Document()
    configure_document(doc, landscape=True, footer_note="NEXT Award 2026 | Inhaltsmaster, nicht offizielle Vorlage | Stand 24.09.2026")
    add_top_rule(doc, MUTED_MINT)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("AMAVIO One")
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(21)
    r.font.color.rgb = RGBColor.from_string(DEEP_INK)
    r2 = p.add_run("  |  Business Model Canvas")
    r2.font.name = "Aptos"
    r2.font.size = Pt(12)
    r2.font.color.rgb = RGBColor.from_string(OP_TEAL)
    sp = doc.add_paragraph("Servicegestütztes Workforce Arrival Management mit Arbeitgeberportal und Partner-/Community-App")
    sp.paragraph_format.space_after = Pt(5)
    sp.runs[0].font.size = Pt(8.5)
    sp.runs[0].font.color.rgb = RGBColor.from_string(OP_TEAL)

    table = doc.add_table(rows=3, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(5.3), Cm(5.3), Cm(5.75), Cm(5.3), Cm(5.3)]
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = widths[idx]

    partners = table.cell(0, 0).merge(table.cell(1, 0))
    value = table.cell(0, 2).merge(table.cell(1, 2))
    segments = table.cell(0, 4).merge(table.cell(1, 4))
    costs = table.cell(2, 0).merge(table.cell(2, 1))
    revenue = table.cell(2, 2).merge(table.cell(2, 4))

    add_bmc_cell(partners, "Schlüsselpartner", [
        "Pilot- und Design-Partner auf Arbeitgeberseite",
        "Technologie, UX, Datenschutz und Security",
        "qualifizierte lokale Service- und Community-Anbieter",
        "Kommunen, Kammern, Welcome-Strukturen und Wirtschaftsförderung",
        "Förder-, Finanzierungs- und später Investmentpartner",
    ])
    add_bmc_cell(table.cell(0, 1), "Schlüsselaktivitäten", [
        "Readiness und Pilotdesign",
        "Arrival-Steuerung, Freigaben und Eskalationen",
        "Portal-, App- und Datenmodellentwicklung",
        "Partnerqualität und Monitoring",
        "Vertrieb und regionale Replikation",
    ], font_size=6.9)
    add_bmc_cell(table.cell(1, 1), "Schlüsselressourcen", [
        "Gründer- und Prozesserfahrung",
        "Playbooks, Status-/Risikomodell und Datenstruktur",
        "Control Center und App als Produkt-IP",
        "lokales Netzwerk und Vertrauen",
        "später Produkt-, Tech- und Customer-Success-Team",
    ], accent=OP_TEAL, font_size=6.9)
    add_bmc_cell(value, "Wertangebote", [
        "geführte Ankunftsstrecke nach der Zusage",
        "Transparenz über Kohorten, Risiken, Aufgaben und Kosten",
        "30-/60-/90-/180-Tage-Monitoring und Lerntransfer",
        "lokale Orientierung, Angebote und Community für Beschäftigte",
        "strukturierter Zugang zu konkreten Bedarfen für Partner",
        "replizierbares System statt fragmentierter Einzelabsprachen",
    ], accent=DARK_TEAL, font_size=7.0)
    add_bmc_cell(table.cell(0, 3), "Kundenbeziehungen", [
        "Readiness Sprint oder begrenzter Pilot",
        "persönliche Steuerung in der frühen Phase",
        "wiederkehrendes Monitoring und Reporting",
        "später Portal- und App-Self-Service",
    ], font_size=6.9)
    add_bmc_cell(table.cell(1, 3), "Kanäle", [
        "direkter Gründervertrieb und Zielaccounts",
        "Empfehlungen und Arbeitgebernetzwerke",
        "Pilot- und Design-Partner",
        "Kammern, Wirtschaftsförderung und kommunale Ökosysteme",
        "Website und partnerschaftlicher Vertrieb",
    ], accent=OP_TEAL, font_size=6.9)
    add_bmc_cell(segments, "Kundensegmente", [
        "Pflegeanbieter, Kliniken und soziale Träger",
        "Hospitality, Infrastruktur und operativer Mittelstand",
        "Entscheider: Geschäftsführung, HR, Standort- und Budgetrollen",
        "sekundär: Beschäftigte sowie lokale Service- und Community-Partner",
    ], font_size=7.0)
    add_bmc_cell(costs, "Kostenstruktur", [
        "Gründer, Team und operative Umsetzung",
        "Produktentwicklung, Hosting, Datenschutz und Security",
        "Vertrieb, Pilotvorbereitung und Customer Success",
        "Recht, Steuer, Versicherung und Buchhaltung",
        "Partneraufbau; externe Fallkosten bleiben getrennt",
    ], accent=CORAL, font_size=6.9)
    add_bmc_cell(revenue, "Einnahmequellen", [
        "Readiness Sprint ab 5.900 EUR netto",
        "Arrival Control Pilot ab 18.500 EUR netto",
        "erweitertes Mandat 42.000 EUR Planpreis bei echter Komplexität",
        "Monitoring ab 2.900 EUR netto pro Monat",
        "Partner Network Build separat; Plattformlizenz erst nach Produktreife",
    ], accent=MUTED_MINT, font_size=7.0)

    doc.save(target)


def main():
    markdown_to_docx(
        ROOT / "01_NEXT_Executive_Summary_Inhaltsmaster_20260924.md",
        ROOT / "AMAVIO_One_NEXT_Executive_Summary_Inhaltsmaster_20260924.docx",
        "NEXT Award 2026 | Inhaltsmaster, vor Einreichung in offizielle Vorlage übertragen | Stand 24.09.2026",
        "Workforce Arrival Management für Arbeitgeber",
    )
    build_bmc(ROOT / "AMAVIO_One_NEXT_Business_Model_Canvas_Inhaltsmaster_20260924.docx")
    markdown_to_docx(
        ROOT / "05_Digital_Tech_to_Product_Go_NoGo_20260924.md",
        ROOT / "AMAVIO_One_Digital_Tech_to_Product_Go_NoGo_20260924.docx",
        "Interne Förderentscheidung | Nicht eingereicht | Stand 24.09.2026",
        "Reife-, Finanzierungs- und Antragsprüfung",
        compact=True,
    )


if __name__ == "__main__":
    main()
