#!/usr/bin/env python3
"""Build the isolated local AMAVIO real-data staging database.

The script prints aggregate counts only. It never logs names, email addresses or
telephone numbers. Apply mode writes a separate SQLite database and never
touches the synthetic CRM state database.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import sqlite3
from datetime import date, datetime, timezone

import openpyxl
from openpyxl.utils.cell import range_boundaries


PROJECT_ROOT = Path(__file__).resolve().parents[1]
CRM_ROOT = PROJECT_ROOT.parent
SOURCE_ROOT = CRM_ROOT / "20261013_AMAVIO_Altdaten_Pruefung" / "01_Arbeitsdaten"
V5_PATH = SOURCE_ROOT / "AMAVIO_Altdaten_Pruefung_und_CRM_Staging_V5_Vollverifikation_20261013.xlsx"
CONFIRMED_PATH = SOURCE_ROOT / "AMAVIO_CRM_Importvorbereitung_Bestaetigt_Aktuell_238_20261013.xlsx"
DRYRUN_PATH = SOURCE_ROOT / "AMAVIO_CRM_R3_10_Feldmapping_und_DryRun_238_20261013.xlsx"
DATA_ROOT = PROJECT_ROOT / "local-data" / "real-data-staging"
DB_PATH = DATA_ROOT / "amavio-real-data-staging-r3-15.sqlite"
SNAPSHOT_ROOT = DATA_ROOT / "snapshots"
QA_ROOT = PROJECT_ROOT / "QA"
DRYRUN_REPORT = QA_ROOT / "CRM_R3_15_REAL_DATA_STAGING_DRY_RUN_20261013.json"
MAPPING_VERSION = "amavio-real-data-staging-r3.15-v1"


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value).strip()


def bool_value(value) -> bool:
    if isinstance(value, bool):
        return value
    return text(value).lower() in {"1", "true", "yes", "ja"}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def stable_id(prefix: str, value: str) -> str:
    return f"{prefix}-{hashlib.sha256(value.encode('utf-8')).hexdigest()[:20]}"


def split_ids(value) -> list[str]:
    raw = text(value)
    if not raw:
        return []
    return [part.strip() for part in re.split(r"[;,|]", raw) if part.strip()]


def table_rows(path: Path, sheet_name: str, table_name: str) -> list[dict]:
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=True)
    worksheet = workbook[sheet_name]
    table = worksheet.tables[table_name]
    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    headers = [text(worksheet.cell(min_row, column).value) for column in range(min_col, max_col + 1)]
    rows = []
    for row_number in range(min_row + 1, max_row + 1):
        values = [worksheet.cell(row_number, column).value for column in range(min_col, max_col + 1)]
        row = dict(zip(headers, values))
        row["__excel_row"] = row_number
        rows.append(row)
    workbook.close()
    return rows


def plain_sheet_rows(path: Path, sheet_name: str) -> list[dict]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name]
    iterator = worksheet.iter_rows(values_only=True)
    headers = [text(value) for value in next(iterator)]
    rows = []
    for row_number, values in enumerate(iterator, start=2):
        row = dict(zip(headers, values))
        row["__excel_row"] = row_number
        rows.append(row)
    workbook.close()
    return rows


def review_bucket(contact_status: str) -> str:
    mapping = {
        "Aktuell bestätigt – keine Kontaktfreigabe": "confirmed_current",
        "Offizieller Stand belegt – Aktualität prüfen": "official_dated",
        "Historisch belegt – nicht aktuell bestätigt": "historical_open",
        "Nicht verwenden – Rollenstand geändert": "changed_legacy",
        "Prüfen – nicht eindeutig aktuell bestätigt": "ambiguous",
    }
    return mapping.get(contact_status, "general_review")


def consent_status(value) -> str:
    candidate = text(value).lower()
    return candidate if candidate in {"blocked", "withdrawn"} else "unknown"


def verified_source_contact_map(rows: list[dict]) -> dict[str, str]:
    mapping = {}
    for row in rows:
        external_id = text(row["External_Contact_ID"])
        for source_id in split_ids(row["Source_Contact_IDs"]):
            if source_id in mapping and mapping[source_id] != external_id:
                raise RuntimeError("Eine Quellkontakt-ID verweist auf mehrere deduplizierte Kontakte.")
            mapping[source_id] = external_id
    return mapping


def load_sources() -> dict:
    hashes = {"v5": sha256(V5_PATH), "confirmed": sha256(CONFIRMED_PATH), "dryrun": sha256(DRYRUN_PATH)}
    return {
        "hashes": hashes,
        "pool_organizations": table_rows(V5_PATH, "Organisationen", "OrganizationsTable"),
        "pool_contacts": table_rows(V5_PATH, "Ansprechpartner", "ContactsTable"),
        "legacy_suppressions": table_rows(V5_PATH, "Sperrliste", "SuppressionTable"),
        "source_register": table_rows(V5_PATH, "Quellenprüfung", "SourceAuditTable"),
        "verified_organizations": plain_sheet_rows(CONFIRMED_PATH, "CRM_Unternehmen"),
        "verified_contacts": plain_sheet_rows(CONFIRMED_PATH, "CRM_Kontakte"),
        "verified_relationships": plain_sheet_rows(CONFIRMED_PATH, "Bestätigte_238"),
        "quarantine": table_rows(DRYRUN_PATH, "Quarantäne", "TblQuarantineR310"),
        "duplicates": table_rows(DRYRUN_PATH, "Dublettenhinweise", "TblDuplicatesR310"),
    }


def validate(data: dict) -> dict:
    pool_orgs = data["pool_organizations"]
    pool_contacts = data["pool_contacts"]
    verified_orgs = data["verified_organizations"]
    verified_contacts = data["verified_contacts"]
    relationships = data["verified_relationships"]

    assertions = {
        "pool_organizations": len(pool_orgs) == 2242,
        "pool_contacts": len(pool_contacts) == 3213,
        "verified_organizations": len(verified_orgs) == 192,
        "verified_contacts": len(verified_contacts) == 166,
        "verified_relationships": len(relationships) == 238,
        "quarantine_items": len(data["quarantine"]) == 93,
        "duplicate_signals": len(data["duplicates"]) == 6,
    }

    pool_org_ids = [text(row["Organisation-ID"]) for row in pool_orgs]
    pool_contact_ids = [text(row["Kontakt-ID"]) for row in pool_contacts]
    verified_org_ids = [text(row["External_Company_ID"]) for row in verified_orgs]
    verified_contact_ids = [text(row["External_Contact_ID"]) for row in verified_contacts]
    assertions.update({
        "pool_org_external_ids_unique": len(set(pool_org_ids)) == len(pool_org_ids) and all(pool_org_ids),
        "pool_contact_external_ids_unique": len(set(pool_contact_ids)) == len(pool_contact_ids) and all(pool_contact_ids),
        "verified_org_external_ids_unique": len(set(verified_org_ids)) == len(verified_org_ids) and all(verified_org_ids),
        "verified_contact_external_ids_unique": len(set(verified_contact_ids)) == len(verified_contact_ids) and all(verified_contact_ids),
        "all_verified_do_not_contact": all(bool_value(row["Do_Not_Contact"]) for row in verified_contacts),
        "all_verified_outreach_blocked": all(not bool_value(row["Outreach_Approved"]) for row in verified_contacts),
    })

    assignment_counts = [len(split_ids(row["Company_IDs"])) for row in verified_contacts]
    multi_counts = [count for count in assignment_counts if count > 1]
    name_review_contacts = sum("prüfen" in text(row["Name_Split_Status"]).lower() for row in verified_contacts)
    name_review_rows = sum("prüfen" in text(row["Name_Split_Status"]).lower() for row in relationships)
    assertions.update({
        "multi_groups_31": len(multi_counts) == 31,
        "multi_assignments_103": sum(multi_counts) == 103,
        "name_review_contacts_48": name_review_contacts == 48,
        "name_review_relationships_73": name_review_rows == 73,
    })

    legacy_buckets = {"official_dated": 0, "historical_open": 0, "changed_legacy": 0, "ambiguous": 0}
    for row in pool_contacts:
        bucket = review_bucket(text(row["Kontaktstatus"]))
        if bucket in legacy_buckets:
            legacy_buckets[bucket] += 1
    assertions["review_legacy_169"] = legacy_buckets == {
        "official_dated": 58,
        "historical_open": 8,
        "changed_legacy": 89,
        "ambiguous": 14,
    }

    source_contact_map = verified_source_contact_map(verified_contacts)
    relationship_contact_ids = {text(row["Contact_ID"]) for row in relationships}
    relationship_org_ids = {text(row["Organisation_ID"]) for row in relationships}
    assertions["relationship_contacts_resolve"] = relationship_contact_ids.issubset(set(source_contact_map))
    assertions["relationship_organizations_resolve"] = relationship_org_ids.issubset(set(verified_org_ids))

    failed = [name for name, passed in assertions.items() if not passed]
    if failed:
        raise RuntimeError(f"Dry-run abgebrochen; fehlgeschlagene Prüfungen: {', '.join(failed)}")

    return {
        "status": "passed",
        "mode": "dry_run_no_write",
        "mappingVersion": MAPPING_VERSION,
        "sourceHashes": data["hashes"],
        "counts": {
            "poolOrganizations": len(pool_orgs),
            "poolContacts": len(pool_contacts),
            "verifiedOrganizations": len(verified_orgs),
            "verifiedContacts": len(verified_contacts),
            "verifiedRelationships": len(relationships),
            "reviewQueue": 238 + sum(legacy_buckets.values()),
            "legacyReview": sum(legacy_buckets.values()),
            "multiGroups": len(multi_counts),
            "multiAssignments": sum(multi_counts),
            "nameReviewContacts": name_review_contacts,
            "nameReviewRelationships": name_review_rows,
            "quarantineItems": len(data["quarantine"]),
            "duplicateSignals": len(data["duplicates"]),
        },
        "reviewBuckets": {"confirmed_current": 238, **legacy_buckets},
        "controls": assertions,
        "outboundEnabled": False,
        "cloudSyncEnabled": False,
        "supabaseEnabled": False,
    }


SCHEMA = """
PRAGMA foreign_keys = ON;
CREATE TABLE IF NOT EXISTS staging_meta (key TEXT PRIMARY KEY, value TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS import_batches (
  id TEXT PRIMARY KEY, mapping_version TEXT NOT NULL, source_hashes_json TEXT NOT NULL,
  mode TEXT NOT NULL, status TEXT NOT NULL, started_at TEXT NOT NULL, completed_at TEXT,
  imported_count INTEGER NOT NULL DEFAULT 0, updated_count INTEGER NOT NULL DEFAULT 0,
  quarantined_count INTEGER NOT NULL DEFAULT 0, outbound_enabled INTEGER NOT NULL DEFAULT 0,
  cloud_sync_enabled INTEGER NOT NULL DEFAULT 0
);
CREATE TABLE IF NOT EXISTS import_rows (
  id TEXT PRIMARY KEY, batch_id TEXT NOT NULL, entity_type TEXT NOT NULL, external_id TEXT NOT NULL,
  outcome TEXT NOT NULL, error_code TEXT, created_at TEXT NOT NULL,
  UNIQUE(batch_id, entity_type, external_id),
  FOREIGN KEY(batch_id) REFERENCES import_batches(id)
);
CREATE TABLE IF NOT EXISTS pool_organizations (
  id TEXT PRIMARY KEY, external_id TEXT NOT NULL UNIQUE, name TEXT NOT NULL, segment TEXT,
  market_priority TEXT, research_usability TEXT, quality_status TEXT, web_check TEXT,
  http_status TEXT, website_source TEXT, current_url TEXT, page_title TEXT, web_checked_at TEXT,
  central_email TEXT, central_phone TEXT, contact_candidates INTEGER NOT NULL DEFAULT 0,
  outreach_approved INTEGER NOT NULL DEFAULT 0, source_reference TEXT, source_note TEXT,
  source_sheet TEXT NOT NULL, source_row INTEGER NOT NULL
);
CREATE TABLE IF NOT EXISTS pool_contacts (
  id TEXT PRIMARY KEY, external_id TEXT NOT NULL UNIQUE, organization_id TEXT, organization_external_id TEXT,
  organization_name TEXT, segment TEXT, market_priority TEXT, priority TEXT, target_role TEXT,
  original_role TEXT, first_name TEXT, last_name TEXT, source_full_name TEXT, name_status TEXT,
  email TEXT, phone TEXT, linkedin_url TEXT, contact_type TEXT, verification_status TEXT,
  evidence_url TEXT, verified_at TEXT, source_contact_status TEXT, consent_status TEXT NOT NULL,
  contact_status TEXT NOT NULL, outreach_eligibility TEXT NOT NULL, do_not_contact INTEGER NOT NULL,
  outreach_approved INTEGER NOT NULL, usage_status TEXT, review_reason TEXT, historical_marketing_status TEXT,
  historical_labels TEXT, review_bucket TEXT NOT NULL, name_review_required INTEGER NOT NULL,
  source_reference TEXT, source_note TEXT, source_sheet TEXT NOT NULL, source_row INTEGER NOT NULL,
  FOREIGN KEY(organization_id) REFERENCES pool_organizations(id)
);
CREATE TABLE IF NOT EXISTS organizations (
  id TEXT PRIMARY KEY, external_id TEXT NOT NULL UNIQUE, display_name TEXT NOT NULL, segment TEXT,
  market_priority TEXT, website TEXT, central_email TEXT, central_phone TEXT,
  verified_assignment_count INTEGER NOT NULL DEFAULT 0, unique_verified_contact_count INTEGER NOT NULL DEFAULT 0,
  outreach_approved INTEGER NOT NULL DEFAULT 0, quality_status TEXT NOT NULL,
  source_reference TEXT, notes TEXT, source_sheet TEXT NOT NULL, source_row INTEGER NOT NULL
);
CREATE TABLE IF NOT EXISTS organization_addresses (
  id TEXT PRIMARY KEY, organization_kind TEXT NOT NULL, organization_id TEXT NOT NULL,
  street TEXT, postal_code TEXT, city TEXT, state TEXT, country TEXT,
  quality_status TEXT NOT NULL, source_sheet TEXT NOT NULL, source_row INTEGER NOT NULL,
  UNIQUE(organization_kind, organization_id)
);
CREATE TABLE IF NOT EXISTS contacts (
  id TEXT PRIMARY KEY, external_id TEXT NOT NULL UNIQUE, first_name TEXT, last_name TEXT,
  source_full_name TEXT, name_status TEXT, role TEXT, email TEXT, phone TEXT, linkedin_url TEXT,
  primary_organization_id TEXT, segment TEXT, verification_status TEXT, verified_at TEXT,
  consent_status TEXT NOT NULL, contact_status TEXT NOT NULL, outreach_eligibility TEXT NOT NULL,
  do_not_contact INTEGER NOT NULL, outreach_approved INTEGER NOT NULL, name_review_required INTEGER NOT NULL,
  evidence_urls TEXT, evidence_type TEXT, source_contact_ids TEXT, source_sheet TEXT NOT NULL,
  source_row INTEGER NOT NULL
);
CREATE TABLE IF NOT EXISTS contact_organization_relationships (
  id TEXT PRIMARY KEY, contact_id TEXT NOT NULL, organization_id TEXT NOT NULL,
  contact_external_id TEXT NOT NULL, organization_external_id TEXT NOT NULL, source_assignment_reference TEXT NOT NULL,
  relationship_type TEXT NOT NULL, is_primary INTEGER NOT NULL, verification_rank TEXT,
  verification_status TEXT, verified_at TEXT, evidence_url TEXT, consent_status TEXT NOT NULL,
  usage_status TEXT, name_review_required INTEGER NOT NULL, source_sheet TEXT NOT NULL,
  source_row INTEGER NOT NULL, UNIQUE(contact_id, organization_id),
  FOREIGN KEY(contact_id) REFERENCES contacts(id), FOREIGN KEY(organization_id) REFERENCES organizations(id)
);
CREATE TABLE IF NOT EXISTS external_references (
  id TEXT PRIMARY KEY, entity_type TEXT NOT NULL, entity_id TEXT NOT NULL, source_system TEXT NOT NULL,
  external_id TEXT NOT NULL, UNIQUE(entity_type, source_system, external_id)
);
CREATE TABLE IF NOT EXISTS provenance (
  id TEXT PRIMARY KEY, entity_type TEXT NOT NULL, entity_id TEXT NOT NULL, source_file TEXT NOT NULL,
  source_sheet TEXT NOT NULL, source_row INTEGER NOT NULL, source_hash TEXT NOT NULL,
  source_reference TEXT, source_note TEXT, imported_at TEXT NOT NULL,
  UNIQUE(entity_type, entity_id, source_file, source_sheet, source_row)
);
CREATE TABLE IF NOT EXISTS review_queue (
  id TEXT PRIMARY KEY, bucket TEXT NOT NULL, rank TEXT, contact_external_id TEXT,
  organization_external_id TEXT, organization_name TEXT, contact_name TEXT, segment TEXT,
  priority TEXT, target_role TEXT, role TEXT, email TEXT, phone TEXT, linkedin_url TEXT,
  verification_status TEXT, verified_at TEXT, evidence_url TEXT, next_action TEXT,
  contact_status TEXT NOT NULL, consent_status TEXT NOT NULL, do_not_contact INTEGER NOT NULL,
  outreach_approved INTEGER NOT NULL, review_reason TEXT, source_sheet TEXT NOT NULL, source_row INTEGER NOT NULL
);
CREATE TABLE IF NOT EXISTS quarantine_items (
  id TEXT PRIMARY KEY, object_type TEXT, source_id TEXT, label TEXT, review_reason TEXT,
  next_decision TEXT, status TEXT, source_sheet TEXT NOT NULL, source_row INTEGER NOT NULL
);
CREATE TABLE IF NOT EXISTS duplicate_signals (
  id TEXT PRIMARY KEY, group_id TEXT, signal_type TEXT, normalized_key TEXT, source_ids TEXT,
  labels TEXT, organization_or_address TEXT, decision_rule TEXT, source_sheet TEXT NOT NULL,
  source_row INTEGER NOT NULL
);
CREATE TABLE IF NOT EXISTS suppression_entries (
  id TEXT PRIMARY KEY, entity_type TEXT NOT NULL, entity_id TEXT NOT NULL, reason_code TEXT NOT NULL,
  reason TEXT NOT NULL, source_status TEXT, consent_status TEXT NOT NULL, effective_at TEXT NOT NULL,
  active INTEGER NOT NULL DEFAULT 1, UNIQUE(entity_type, entity_id, reason_code)
);
CREATE TABLE IF NOT EXISTS legacy_suppressions (
  id TEXT PRIMARY KEY, email TEXT, organization_name TEXT, contact_name TEXT, suppression_status TEXT,
  reason TEXT, source_reference TEXT, rule_text TEXT, source_sheet TEXT NOT NULL, source_row INTEGER NOT NULL
);
CREATE TABLE IF NOT EXISTS source_register (
  id TEXT PRIMARY KEY, source_file TEXT, source_sheet_name TEXT, data_rows INTEGER,
  amavio_usage TEXT, risk_boundary TEXT, decision TEXT, staging_source_sheet TEXT NOT NULL,
  staging_source_row INTEGER NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_pool_org_segment ON pool_organizations(segment);
CREATE INDEX IF NOT EXISTS idx_pool_org_quality ON pool_organizations(quality_status);
CREATE INDEX IF NOT EXISTS idx_pool_contact_segment ON pool_contacts(segment);
CREATE INDEX IF NOT EXISTS idx_pool_contact_bucket ON pool_contacts(review_bucket);
CREATE INDEX IF NOT EXISTS idx_pool_contact_consent ON pool_contacts(consent_status);
CREATE INDEX IF NOT EXISTS idx_review_bucket ON review_queue(bucket);
CREATE INDEX IF NOT EXISTS idx_relationship_contact ON contact_organization_relationships(contact_id);
CREATE INDEX IF NOT EXISTS idx_relationship_org ON contact_organization_relationships(organization_id);
CREATE INDEX IF NOT EXISTS idx_provenance_entity ON provenance(entity_type, entity_id);
"""


def upsert_many(connection: sqlite3.Connection, table: str, columns: list[str], rows: list[tuple]) -> tuple[int, int]:
    if not rows:
        return 0, 0
    key_index = columns.index("id")
    existing = {row[0] for row in connection.execute(f"SELECT id FROM {table}")}
    incoming = {row[key_index] for row in rows}
    created = len(incoming - existing)
    updated = len(incoming & existing)
    placeholders = ", ".join("?" for _ in columns)
    updates = ", ".join(f"{column}=excluded.{column}" for column in columns if column != "id")
    connection.executemany(
        f"INSERT INTO {table} ({', '.join(columns)}) VALUES ({placeholders}) ON CONFLICT(id) DO UPDATE SET {updates}",
        rows,
    )
    return created, updated


def apply_import(data: dict, report: dict) -> dict:
    DATA_ROOT.mkdir(parents=True, exist_ok=True, mode=0o700)
    SNAPSHOT_ROOT.mkdir(parents=True, exist_ok=True, mode=0o700)
    QA_ROOT.mkdir(parents=True, exist_ok=True)
    os.chmod(DATA_ROOT, 0o700)
    os.chmod(SNAPSHOT_ROOT, 0o700)

    snapshot_path = None
    if DB_PATH.exists() and DB_PATH.stat().st_size:
        stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
        snapshot_path = SNAPSHOT_ROOT / f"amavio-real-data-staging-before-{stamp}.sqlite"
        shutil.copy2(DB_PATH, snapshot_path)
        os.chmod(snapshot_path, 0o600)

    now = datetime.now(timezone.utc).isoformat()
    batch_id = f"batch-{hashlib.sha256((MAPPING_VERSION + json.dumps(data['hashes'], sort_keys=True)).encode()).hexdigest()[:20]}"
    connection = sqlite3.connect(DB_PATH)
    connection.execute("PRAGMA journal_mode=WAL")
    connection.executescript(SCHEMA)
    created_total = 0
    updated_total = 0

    def account(result: tuple[int, int]):
        nonlocal created_total, updated_total
        created_total += result[0]
        updated_total += result[1]

    try:
        connection.execute("BEGIN IMMEDIATE")
        connection.execute(
            "INSERT INTO import_batches(id, mapping_version, source_hashes_json, mode, status, started_at, outbound_enabled, cloud_sync_enabled) VALUES (?, ?, ?, ?, ?, ?, 0, 0) ON CONFLICT(id) DO UPDATE SET status=excluded.status, started_at=excluded.started_at, completed_at=NULL",
            (batch_id, MAPPING_VERSION, json.dumps(data["hashes"], sort_keys=True), "local_real_data_staging", "running", now),
        )
        meta = {
            "mode": "real_data_staging",
            "maturity": "lokales Echtdaten-Staging – keine Kontaktfreigabe",
            "outboundEnabled": "false",
            "cloudSyncEnabled": "false",
            "supabaseEnabled": "false",
            "appIntegrationEnabled": "false",
            "accIntegrationEnabled": "false",
            "mappingVersion": MAPPING_VERSION,
            "updatedAt": now,
        }
        connection.executemany("INSERT INTO staging_meta(key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", list(meta.items()))

        pool_org_rows = []
        pool_addr_rows = []
        external_rows = []
        provenance_rows = []
        import_rows = []
        pool_org_id_by_external = {}
        for row in data["pool_organizations"]:
            external_id = text(row["Organisation-ID"])
            entity_id = stable_id("pool-org", external_id)
            pool_org_id_by_external[external_id] = entity_id
            pool_org_rows.append((
                entity_id, external_id, text(row["Name"]), text(row["Segment"]), text(row["Marktpriorität"]),
                text(row["Research-Nutzbarkeit"]), text(row["Qualitätsstatus"]), text(row["Web-Check"]), text(row["HTTP"]),
                text(row["Website Quelle"]), text(row["Aktuelle Ziel-URL"]), text(row["Seitentitel"]), text(row["Web-Prüfdatum"]),
                text(row["Zentrale E-Mail"]), text(row["Zentrale Telefon"]), int(row["Kontaktkandidaten"] or 0), 0,
                text(row["Quellen"]), text(row["Quellenhinweis"]), "Organisationen", int(row["__excel_row"]),
            ))
            address_id = stable_id("pool-addr", external_id)
            pool_addr_rows.append((address_id, "pool", entity_id, text(row["Adresse"]), text(row["PLZ"]), text(row["Ort"]), text(row["Bundesland"]), text(row["Land"]), text(row["Qualitätsstatus"]), "Organisationen", int(row["__excel_row"])))
            external_rows.append((stable_id("xref", f"pool-org|{external_id}"), "pool_organization", entity_id, "v5_legacy_staging", external_id))
            provenance_rows.append((stable_id("prov", f"pool-org|{external_id}"), "pool_organization", entity_id, V5_PATH.name, "Organisationen", int(row["__excel_row"]), data["hashes"]["v5"], text(row["Quellen"]), text(row["Quellenhinweis"]), now))
            import_rows.append((stable_id("import-row", f"{batch_id}|pool-org|{external_id}"), batch_id, "pool_organization", external_id, "imported", None, now))
        account(upsert_many(connection, "pool_organizations", ["id", "external_id", "name", "segment", "market_priority", "research_usability", "quality_status", "web_check", "http_status", "website_source", "current_url", "page_title", "web_checked_at", "central_email", "central_phone", "contact_candidates", "outreach_approved", "source_reference", "source_note", "source_sheet", "source_row"], pool_org_rows))
        account(upsert_many(connection, "organization_addresses", ["id", "organization_kind", "organization_id", "street", "postal_code", "city", "state", "country", "quality_status", "source_sheet", "source_row"], pool_addr_rows))

        pool_contact_rows = []
        suppression_rows = []
        for row in data["pool_contacts"]:
            external_id = text(row["Kontakt-ID"])
            entity_id = stable_id("pool-contact", external_id)
            organization_external_id = text(row["Organisation-ID"])
            organization_id = pool_org_id_by_external.get(organization_external_id)
            source_consent = consent_status(row["Consent-Status"])
            source_status = text(row["Kontaktstatus"])
            bucket = review_bucket(source_status)
            name_status = text(row["Namenstatus"])
            name_review = int("prüfen" in name_status.lower() or name_status in {"Name fehlt", "Ansprechpartner nicht benannt", "Funktionsbezeichnung, keine Person"})
            pool_contact_rows.append((
                entity_id, external_id, organization_id, organization_external_id, text(row["Organisation"]), text(row["Segment"]),
                text(row["Marktpriorität"]), text(row["Priorität"]), text(row["Zielrolle"]), text(row["Rolle original"]),
                text(row["Vorname"]), text(row["Nachname"]), text(row["Vollname original"]), name_status, text(row["E-Mail"]),
                text(row["Telefon"]), text(row["LinkedIn"]), text(row["Kontaktart"]), text(row["Kontaktverifikation"]),
                text(row["Aktuelle Kontaktquelle"]), text(row["Prüfdatum"]), source_status, source_consent, "do_not_contact",
                "blocked", 1, 0, text(row["Nutzungsstatus"]), text(row["Sperr-/Prüfgrund"]),
                text(row["Historischer Marketingstatus"]), text(row["Historische Labels"]), bucket, name_review,
                text(row["Quellen"]), text(row["Quellenhinweis"]), "Ansprechpartner", int(row["__excel_row"]),
            ))
            external_rows.append((stable_id("xref", f"pool-contact|{external_id}"), "pool_contact", entity_id, "v5_legacy_staging", external_id))
            provenance_rows.append((stable_id("prov", f"pool-contact|{external_id}"), "pool_contact", entity_id, V5_PATH.name, "Ansprechpartner", int(row["__excel_row"]), data["hashes"]["v5"], text(row["Quellen"]), text(row["Quellenhinweis"]), now))
            reason = text(row["Sperr-/Prüfgrund"]) or "Keine Außenkontaktfreigabe"
            suppression_rows.append((stable_id("suppress", f"pool-contact|{external_id}"), "pool_contact", entity_id, "do_not_contact", reason, text(row["Nutzungsstatus"]), source_consent, now, 1))
            import_rows.append((stable_id("import-row", f"{batch_id}|pool-contact|{external_id}"), batch_id, "pool_contact", external_id, "review_required" if name_review else "imported", "NAME_REVIEW_REQUIRED" if name_review else None, now))
        account(upsert_many(connection, "pool_contacts", ["id", "external_id", "organization_id", "organization_external_id", "organization_name", "segment", "market_priority", "priority", "target_role", "original_role", "first_name", "last_name", "source_full_name", "name_status", "email", "phone", "linkedin_url", "contact_type", "verification_status", "evidence_url", "verified_at", "source_contact_status", "consent_status", "contact_status", "outreach_eligibility", "do_not_contact", "outreach_approved", "usage_status", "review_reason", "historical_marketing_status", "historical_labels", "review_bucket", "name_review_required", "source_reference", "source_note", "source_sheet", "source_row"], pool_contact_rows))

        verified_org_rows = []
        verified_org_addr_rows = []
        verified_org_id_by_external = {}
        for row in data["verified_organizations"]:
            external_id = text(row["External_Company_ID"])
            entity_id = stable_id("verified-org", external_id)
            verified_org_id_by_external[external_id] = entity_id
            address_complete = all(text(row[field]) for field in ("Street", "Postal_Code", "City"))
            verified_org_rows.append((entity_id, external_id, text(row["Company_Name"]), text(row["Segment"]), text(row["Market_Priority"]), text(row["Website"]), text(row["Central_Email"]), text(row["Central_Phone"]), int(row["Verified_Assignment_Count"] or 0), int(row["Unique_Verified_Contact_Count"] or 0), 0, "confirmed_current" if address_complete else "address_review_required", text(row["Sources"]), text(row["Organisation_Note"]), "CRM_Unternehmen", int(row["__excel_row"])))
            verified_org_addr_rows.append((stable_id("verified-addr", external_id), "verified", entity_id, text(row["Street"]), text(row["Postal_Code"]), text(row["City"]), text(row["State"]), text(row["Country"]), "confirmed" if address_complete else "review_required", "CRM_Unternehmen", int(row["__excel_row"])))
            external_rows.append((stable_id("xref", f"verified-org|{external_id}"), "organization", entity_id, "confirmed_import_20261013", external_id))
            provenance_rows.append((stable_id("prov", f"verified-org|{external_id}"), "organization", entity_id, CONFIRMED_PATH.name, "CRM_Unternehmen", int(row["__excel_row"]), data["hashes"]["confirmed"], text(row["Sources"]), text(row["Organisation_Note"]), now))
            import_rows.append((stable_id("import-row", f"{batch_id}|verified-org|{external_id}"), batch_id, "organization", external_id, "review_required" if not address_complete else "imported", "ADDRESS_REVIEW_REQUIRED" if not address_complete else None, now))
        account(upsert_many(connection, "organizations", ["id", "external_id", "display_name", "segment", "market_priority", "website", "central_email", "central_phone", "verified_assignment_count", "unique_verified_contact_count", "outreach_approved", "quality_status", "source_reference", "notes", "source_sheet", "source_row"], verified_org_rows))
        account(upsert_many(connection, "organization_addresses", ["id", "organization_kind", "organization_id", "street", "postal_code", "city", "state", "country", "quality_status", "source_sheet", "source_row"], verified_org_addr_rows))

        verified_contact_rows = []
        verified_contact_id_by_external = {}
        verified_contact_company_ids = {}
        source_contact_map = verified_source_contact_map(data["verified_contacts"])
        for row in data["verified_contacts"]:
            external_id = text(row["External_Contact_ID"])
            entity_id = stable_id("verified-contact", external_id)
            verified_contact_id_by_external[external_id] = entity_id
            company_ids = split_ids(row["Company_IDs"])
            verified_contact_company_ids[external_id] = company_ids
            primary_org_id = verified_org_id_by_external.get(company_ids[0]) if len(company_ids) == 1 else None
            name_review = int("prüfen" in text(row["Name_Split_Status"]).lower())
            source_consent = consent_status(row["Consent_Status"])
            verified_contact_rows.append((entity_id, external_id, text(row["First_Name"]), text(row["Last_Name"]), text(row["Full_Name"]), text(row["Name_Split_Status"]), text(row["Job_Title"]), text(row["Email"]), text(row["Phone"]), text(row["LinkedIn_URL"]), primary_org_id, text(row["Segment"]), text(row["Verification_Status"]), text(row["Verification_Date"]), source_consent, "do_not_contact", "blocked", 1, 0, name_review, text(row["Evidence_URLs"]), text(row["Evidence_Type"]), text(row["Source_Contact_IDs"]), "CRM_Kontakte", int(row["__excel_row"])))
            external_rows.append((stable_id("xref", f"verified-contact|{external_id}"), "contact", entity_id, "confirmed_import_20261013", external_id))
            provenance_rows.append((stable_id("prov", f"verified-contact|{external_id}"), "contact", entity_id, CONFIRMED_PATH.name, "CRM_Kontakte", int(row["__excel_row"]), data["hashes"]["confirmed"], text(row["Evidence_URLs"]), text(row["Source_Contact_IDs"]), now))
            suppression_rows.append((stable_id("suppress", f"verified-contact|{external_id}"), "contact", entity_id, "do_not_contact", "Keine Außenkontaktfreigabe", "BLOCKIERT – keine Außenkontaktfreigabe", source_consent, now, 1))
            import_rows.append((stable_id("import-row", f"{batch_id}|verified-contact|{external_id}"), batch_id, "contact", external_id, "review_required" if name_review else "imported", "NAME_REVIEW_REQUIRED" if name_review else None, now))
        account(upsert_many(connection, "contacts", ["id", "external_id", "first_name", "last_name", "source_full_name", "name_status", "role", "email", "phone", "linkedin_url", "primary_organization_id", "segment", "verification_status", "verified_at", "consent_status", "contact_status", "outreach_eligibility", "do_not_contact", "outreach_approved", "name_review_required", "evidence_urls", "evidence_type", "source_contact_ids", "source_sheet", "source_row"], verified_contact_rows))

        relationship_rows = []
        relationship_count_by_contact = {}
        for row in data["verified_relationships"]:
            contact_external_id = source_contact_map[text(row["Contact_ID"])]
            relationship_count_by_contact[contact_external_id] = relationship_count_by_contact.get(contact_external_id, 0) + 1
        for row in data["verified_relationships"]:
            source_assignment_reference = text(row["Contact_ID"])
            contact_external_id = source_contact_map[source_assignment_reference]
            organization_external_id = text(row["Organisation_ID"])
            contact_id = verified_contact_id_by_external[contact_external_id]
            organization_id = verified_org_id_by_external[organization_external_id]
            relationship_id = stable_id("relationship", f"{contact_external_id}|{organization_external_id}")
            is_primary = int(relationship_count_by_contact[contact_external_id] == 1)
            name_review = int("prüfen" in text(row["Name_Split_Status"]).lower())
            relationship_rows.append((relationship_id, contact_id, organization_id, contact_external_id, organization_external_id, source_assignment_reference, "primary" if is_primary else "additional_unresolved", is_primary, text(row["Verification_Rank"]), text(row["Verification_Status"]), text(row["Verification_Date"]), text(row["Evidence_URL"]), consent_status(row["Consent_Status"]), text(row["Usage_Status"]), name_review, "Bestätigte_238", int(row["__excel_row"])))
            provenance_rows.append((stable_id("prov", f"relationship|{contact_external_id}|{organization_external_id}"), "contact_organization_relationship", relationship_id, CONFIRMED_PATH.name, "Bestätigte_238", int(row["__excel_row"]), data["hashes"]["confirmed"], text(row["Evidence_URL"]), text(row["Source_Note"]), now))
            import_rows.append((stable_id("import-row", f"{batch_id}|relationship|{contact_external_id}|{organization_external_id}"), batch_id, "contact_organization_relationship", relationship_id, "review_required" if name_review else "imported", "NAME_REVIEW_REQUIRED" if name_review else None, now))
        account(upsert_many(connection, "contact_organization_relationships", ["id", "contact_id", "organization_id", "contact_external_id", "organization_external_id", "source_assignment_reference", "relationship_type", "is_primary", "verification_rank", "verification_status", "verified_at", "evidence_url", "consent_status", "usage_status", "name_review_required", "source_sheet", "source_row"], relationship_rows))

        review_rows = []
        for row in data["verified_relationships"]:
            contact_external_id = source_contact_map[text(row["Contact_ID"])]
            organization_external_id = text(row["Organisation_ID"])
            queue_id = stable_id("review", f"current|{contact_external_id}|{organization_external_id}")
            review_rows.append((queue_id, "confirmed_current", text(row["Verification_Rank"]), contact_external_id, organization_external_id, text(row["Organisation"]), text(row["Full_Name"]), text(row["Segment"]), text(row["Priority"]), text(row["Target_Role"]), text(row["Current_Role"]), text(row["Current_Email"]), text(row["Current_Phone"]), text(row["LinkedIn_URL"]), text(row["Verification_Status"]), text(row["Verification_Date"]), text(row["Evidence_URL"]), "Im lokalen Staging prüfen; keine Kontaktfreigabe", "do_not_contact", consent_status(row["Consent_Status"]), 1, 0, text(row["Suppression_Reason"]), "Bestätigte_238", int(row["__excel_row"])))
        for row in data["pool_contacts"]:
            bucket = review_bucket(text(row["Kontaktstatus"]))
            if bucket not in {"official_dated", "historical_open", "changed_legacy", "ambiguous"}:
                continue
            external_id = text(row["Kontakt-ID"])
            queue_id = stable_id("review", f"legacy|{external_id}")
            review_rows.append((queue_id, bucket, "", external_id, text(row["Organisation-ID"]), text(row["Organisation"]), text(row["Vollname original"]), text(row["Segment"]), text(row["Priorität"]), text(row["Zielrolle"]), text(row["Rolle original"]), text(row["E-Mail"]), text(row["Telefon"]), text(row["LinkedIn"]), text(row["Kontaktverifikation"]), text(row["Prüfdatum"]), text(row["Aktuelle Kontaktquelle"]), "Manuell prüfen; nicht als aktueller Lead verwenden", "do_not_contact", consent_status(row["Consent-Status"]), 1, 0, text(row["Sperr-/Prüfgrund"]), "Ansprechpartner", int(row["__excel_row"])))
        account(upsert_many(connection, "review_queue", ["id", "bucket", "rank", "contact_external_id", "organization_external_id", "organization_name", "contact_name", "segment", "priority", "target_role", "role", "email", "phone", "linkedin_url", "verification_status", "verified_at", "evidence_url", "next_action", "contact_status", "consent_status", "do_not_contact", "outreach_approved", "review_reason", "source_sheet", "source_row"], review_rows))

        quarantine_rows = []
        for row in data["quarantine"]:
            source_id = text(row["Source_ID"])
            item_id = stable_id("quarantine", f"{text(row['Objekt'])}|{source_id}|{text(row['Prüfgrund'])}")
            quarantine_rows.append((item_id, text(row["Objekt"]), source_id, text(row["Bezeichnung"]), text(row["Prüfgrund"]), text(row["Nächste Entscheidung"]), text(row["Status"]), "Quarantäne", int(row["__excel_row"])))
            import_rows.append((stable_id("import-row", f"{batch_id}|quarantine|{item_id}"), batch_id, "quarantine_item", item_id, "quarantined", text(row["Prüfgrund"]), now))
        account(upsert_many(connection, "quarantine_items", ["id", "object_type", "source_id", "label", "review_reason", "next_decision", "status", "source_sheet", "source_row"], quarantine_rows))

        duplicate_rows = []
        for row in data["duplicates"]:
            item_id = stable_id("duplicate", f"{text(row['Gruppe'])}|{text(row['Typ'])}|{text(row['Source_IDs'])}")
            duplicate_rows.append((item_id, text(row["Gruppe"]), text(row["Typ"]), text(row["Normalisierter_Schlüssel"]), text(row["Source_IDs"]), text(row["Namen"]), text(row["Organisationen_oder_Adressen"]), text(row["Entscheidungsregel"]), "Dublettenhinweise", int(row["__excel_row"])))
        account(upsert_many(connection, "duplicate_signals", ["id", "group_id", "signal_type", "normalized_key", "source_ids", "labels", "organization_or_address", "decision_rule", "source_sheet", "source_row"], duplicate_rows))

        legacy_suppression_rows = []
        for row in data["legacy_suppressions"]:
            item_id = stable_id("legacy-suppress", f"{text(row['E-Mail'])}|{text(row['Organisation'])}|{text(row['Grund'])}")
            legacy_suppression_rows.append((item_id, text(row["E-Mail"]), text(row["Organisation"]), text(row["Kontakt"]), text(row["Sperrstatus"]), text(row["Grund"]), text(row["Quelle"]), text(row["Regel"]), "Sperrliste", int(row["__excel_row"])))
        account(upsert_many(connection, "legacy_suppressions", ["id", "email", "organization_name", "contact_name", "suppression_status", "reason", "source_reference", "rule_text", "source_sheet", "source_row"], legacy_suppression_rows))

        source_rows = []
        for row in data["source_register"]:
            item_id = stable_id("source", f"{text(row['Datei'])}|{text(row['Blatt'])}")
            source_rows.append((item_id, text(row["Datei"]), text(row["Blatt"]), int(row["Datenzeilen"] or 0), text(row["Nutzung für AMAVIO"]), text(row["Risiko / Grenze"]), text(row["Entscheidung"]), "Quellenprüfung", int(row["__excel_row"])))
        account(upsert_many(connection, "source_register", ["id", "source_file", "source_sheet_name", "data_rows", "amavio_usage", "risk_boundary", "decision", "staging_source_sheet", "staging_source_row"], source_rows))

        account(upsert_many(connection, "external_references", ["id", "entity_type", "entity_id", "source_system", "external_id"], external_rows))
        account(upsert_many(connection, "provenance", ["id", "entity_type", "entity_id", "source_file", "source_sheet", "source_row", "source_hash", "source_reference", "source_note", "imported_at"], provenance_rows))
        account(upsert_many(connection, "suppression_entries", ["id", "entity_type", "entity_id", "reason_code", "reason", "source_status", "consent_status", "effective_at", "active"], suppression_rows))
        account(upsert_many(connection, "import_rows", ["id", "batch_id", "entity_type", "external_id", "outcome", "error_code", "created_at"], import_rows))

        quarantined = connection.execute("SELECT COUNT(*) FROM import_rows WHERE batch_id=? AND outcome='quarantined'", (batch_id,)).fetchone()[0]
        connection.execute("UPDATE import_batches SET status='completed', completed_at=?, imported_count=?, updated_count=?, quarantined_count=? WHERE id=?", (now, created_total, updated_total, quarantined, batch_id))
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()

    os.chmod(DB_PATH, 0o600)
    for suffix in ("-wal", "-shm"):
        sidecar = Path(f"{DB_PATH}{suffix}")
        if sidecar.exists():
            os.chmod(sidecar, 0o600)

    return {
        "status": "completed",
        "database": str(DB_PATH.relative_to(PROJECT_ROOT)),
        "batchId": batch_id,
        "created": created_total,
        "updated": updated_total,
        "quarantined": report["counts"]["quarantineItems"],
        "snapshotCreated": bool(snapshot_path),
        "snapshot": str(snapshot_path.relative_to(PROJECT_ROOT)) if snapshot_path else None,
        "outboundEnabled": False,
        "cloudSyncEnabled": False,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--apply", action="store_true", help="Write the isolated local staging SQLite database")
    arguments = parser.parse_args()

    data = load_sources()
    report = validate(data)
    QA_ROOT.mkdir(parents=True, exist_ok=True)
    DRYRUN_REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    output = {"dryRun": report}
    if arguments.apply:
        output["apply"] = apply_import(data, report)
    print(json.dumps(output, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
